#!/usr/bin/env python3
"""DINO-X report from the Mac mirror of the Sol run tree: numbers workbook, PNG figures,
a graphs-only 16:9 deck and the HTML tables for the PDF ledger. Every number is read from a file under
--root; an unreadable value leaves an empty cell and a warning, never a guess.

  <python> -B plot/build_step8_report.py [--root DIR] [--container DIR] [--out-dir DIR] [--skip-figs] [--selftest]
"""

import argparse
import contextlib
import csv
import html
import io
import json
import math
import os
import re
import shutil
import statistics
import subprocess
import sys
import tempfile
import time
import traceback
import types
from collections import Counter
from datetime import datetime
from pathlib import Path

sys.dont_write_bytecode = True  # importing plot_linear_probes must not drop a __pycache__ into the plot folder
SCRATCH = Path("/private/tmp/claude-501/-Users-jaroddefilippo-Jarod-s-Second-Brain/2a683ce5-341a-44cb-a02d-a29f78df69f6/scratchpad")
os.environ["MPLCONFIGDIR"] = str(SCRATCH / "mplconfig")  # matplotlib's cache stays in the scratchpad (subprocesses inherit it)
PLOT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(PLOT_DIR))
from run_labels import arm_label  # noqa: E402  (the fixed labels of the input-size arms, shared with compare_cyclic_runs.py)

ROOT = Path("/Users/jaroddefilippo/code/CSE494-IP/demo1_results")
CONTAINER = Path("/Users/jaroddefilippo/code/CSE494-IP/2026Fall-CSE507-Ark-Container-main")
XLSX_NAME, PPTX_NAME = "DINO-X-numbers.xlsx", "DINO-X-graphs.pptx"  # the project is DINO-X; until 2026-10-18 these were step8-dinov3-numbers.xlsx and step8-dinov3-graphs.pptx (same data)
DATASETS = ("VinDrCXR", "ChestXray14", "MIMIC", "CheXpert")  # training order inside a cycle
SHORT = dict(zip(DATASETS, ("VinDr", "CXR14", "MIMIC", "CheXpert")))
TITLE_DS = dict(zip(DATASETS, ("VinDr-CXR", "ChestX-ray14", "MIMIC-CXR", "CheXpert")))  # display names: titles, legends, tables; the ids stay in DATASETS, SHORT and the workbook
HEAD_DS = TITLE_DS  # the dataset names in the generated HTML table headers (the workbook column ids keep SHORT)
POINTS = tuple("after_" + d for d in DATASETS)
END = "after_CheXpert"  # last evaluation of a cycle = END of the cycle: the ledger's headline numbers
INIT = "init"  # cycle 0: the initial model before any training; the student and the teacher each have a value of their own (two separate model builds). Both the end-of-cycle and the focused value of every dataset. Found by cycle / evaluation_point, never by epoch (cycle 1 is epoch 0 too)
PROTOCOLS = ("tencrop", "centercrop")  # the two offline evaluations of the saved weights: cycle_<NNNN>_<run>.csv and cycle_<NNNN>_<run>.centercrop.csv at the mirror root
CC_NOTE = "hollow = offline evaluation of the saved weights"  # the one legend note of a center-crop figure that has offline points
SPLIT_FILES = (("test", "test_performance.csv"), ("validation", "val_performance.csv"))
N_FILES = len(SPLIT_FILES) * len(DATASETS)

FIXED_RUNS = (
    ("A1", "all4_concurrent_random_sampling_dinov3_seed_100", "A1 (SGD 0.01, early arm)", "early"),
    ("A2", "all4_cyclic_dinov3_v2_seed_100", "A2 (SGD 0.0025, early arm)", "early"),
    ("B", "all4_cyclic_dinov3_adamw_seed_100", "B (first AdamW arm)", "early"),
    ("C", "all4_cyclic_dinov3_short_seed_100", "C (short cosine)", "early"),
    ("D", "all4_cyclic_dinov3_frozen_seed_100", "D (frozen, 5 cycles)", "early"),
    ("D-long", "all4_cyclic_dinov3_frozen_long_seed_100", "D-long (frozen control, 50 cycles)", "early"),
    ("probe_a", "all4_cyclic_dinov3_probe_a_seed_100", "probe a (pilot a e4, frozen 5 cycles)", "gate"),
    ("probe_b", "all4_cyclic_dinov3_probe_b_seed_100", "probe b (pilot b e4, frozen 5 cycles)", "gate"),
)
RECIPE_KEYS = ("opt", "lr", "wd", "layer_decay", "patch_embed_mult", "warmup", "cycles")
RECIPE_ABBR = {"layer_decay": "ld", "patch_embed_mult": "pe", "init_tag": "init"}
CENTER = "S2"
PROBE_TAGS = (
    ("rad_dino", "RAD-DINO ViT-B/14 (Microsoft)"),
    ("raydino", "RayDINO ViT-L/16 (Meta, XRay-DINO release)"),
    ("dinov2", "DINOv2 ViT-B/14 (Meta, LVD-142M)"),
    ("stock_dinov3", "DINOv3 ViT-B/16 stock (Meta, LVD-1689M)"),
    ("stageA_e0010", "DINO-X Cosine e10"),
    ("stageA_e0020", "DINO-X Cosine e20"),
    ("stageA_final", "DINO-X Cosine e30"),
    ("constlr_e0010", "DINO-X ConstLR e10"),
    ("constlr_e0020", "DINO-X ConstLR e20"),
    ("constlr_e0030", "DINO-X ConstLR e30"),
)
# DINO-X = our DINOv3 ViT-B/16 continued-pretrained on the four chest X-ray datasets; two pretraining runs: tag main (cosine schedule) and tag constlr (constant LR).
SSL_NAMES = {"main": "DINO-X Cosine", "constlr": "DINO-X ConstLR"}
CONSTLR = re.compile(r"constlr_e(\d{4})")  # the constant-LR run's probed exports, results/constlr_e<epoch>/: every such directory is admitted, listed in PROBE_TAGS or not
PROBE_PAGE_TAGS = ("rad_dino", "raydino", "dinov2", "stock_dinov3", "stageA_final")  # probe pages: the references and the cosine run's final export (e30); the constant-LR run's latest is picked by probe_page_tags
# Published Ark+ reference line, quoted by the spec (Ma et al., Nature 2025, Fig. 3a); the one number in this file that is not read from the mirror.
ARK_REF = ("VinDrCXR", "Ark+ Swin-L linear probe, 6 global labels (Ma et al., Nature 2025, Fig. 3a)", "96.19")
EARLY = ("A1", "A2", "B", "C", "D", "D-long")
SWEEP = ("S2", "S1", "S3", "S4", "S5", "S6", "S7", "S13", "S16", "S10", "S11", "S12")
DECK_ORDER = (*SWEEP, "D-long", "B", "C", "D", "A1", "A2", "probe_a", "probe_b")
TAB20 = ("#1f77b4", "#aec7e8", "#ff7f0e", "#ffbb78", "#2ca02c", "#98df8a", "#d62728", "#ff9896", "#9467bd", "#c5b0d5", "#8c564b", "#c49c94", "#e377c2", "#f7b6d2", "#7f7f7f",
         "#c7c7c7", "#bcbd22", "#dbdb8d", "#17becf", "#9edae5")  # matplotlib's tab20
# One colour per arm, the same in every figure that shows the arm: tab20 entries 0..14 in SWEEP order, then the early arms B, C and D-long. An arm of sweep/wave1.tsv that this palette lacks
# (an id SWEEP does not name, e.g. S2b) takes the next of EXTRA_COLORS in the table's file order (run_table), counted over every table row so its colour does not wait for earlier arms' runs to exist;
# the file order does nothing else, the workbook and the deck follow the fixed presentation order. EXTRA_COLORS skips the grays: tab20 entries 16, 17, 18, 19, then the pale gray 15 last (14 is D-long's).
# ponytail: five colours; a sixth extra arm starts the sequence again. A longer palette if that ever happens.
ARM_COLORS = dict(zip((*SWEEP, "B", "C", "D-long"), TAB20))
EXTRA_COLORS = (*TAB20[16:], TAB20[15])
SEED_COLS = ("seed", "eval_every", "eval_cycle0")  # the table's trailing columns; a table without them is a seed-100 table, the other two empty
# The input-size arms S18, S19 (wave1.tsv's input_size column; S2 is the 224 px control): their run label and title are run_labels.ARM_LABELS' fixed strings
# (arm_label). Their ids are not in SWEEP, so they take the extra arms' place (after S16) and colours (tab20 16, 17 in the table's order)
COMPARE = {
    "adamw": ("AdamW arms", ("S2", "S1", "S3", "S4", "S5", "S6", "S7", "S13", "S16")),
    "init": ("initialisation arms", ("S2", "S13", "S16")),
    "sgd_frozen": ("SGD arms and frozen control", ("S2", "S10", "S11", "S12", "D-long")),
}
FIG_KINDS = (("focused_unfocused", "test"), ("focused_unfocused", "validation"), ("focused", "test"), ("unfocused", "test"))
SSL_RUNS = ("main", "constlr", "pilot_a", "pilot_b")  # "timing" is a throughput probe, not a training run
LOSS_COLS = ("dino_global", "dino_local", "koleo", "ibot", "gram", "total")
SSL_HEADER = ["run", "epoch", "steps", *LOSS_COLS, "lr_end", "teacher_temp_end", "step_time_mean"]
LOSS_FIGS = (("dino_global", "DINO-X pretraining loss — dino_global"), ("ibot", "DINO-X pretraining loss — ibot"), ("koleo", "DINO-X pretraining loss — koleo"),
             ("gram", "DINO-X pretraining loss — gram"), ("lr", "DINO-X learning rate (end of epoch)"))
END_COLS = [f"{w}_{SHORT[d]}_end" for w in "TS" for d in DATASETS]
FOC_COLS = [f"{w}_{SHORT[d]}_focused" for w in "TS" for d in DATASETS]
AUC_FMT, LOSS_FMT = "0.00", "0.0000"
WARNINGS, BAD, MANIFEST = [], Counter(), {"inputs": {}, "outputs": []}


def warn(msg):
    WARNINGS.append(msg)


def num(s):
    """CSV text to float; blank, nan and inf (nan AUC_Fracture on the CheXpert validation list) give None."""
    try:
        v = float(s)
    except (TypeError, ValueError):
        return None
    return v if math.isfinite(v) else None


def pct(v):
    return None if v is None else round(v * 100, 4)


def f1(v):
    return "" if v is None else f"{v * 100:.1f}"


def p1(v):
    """The number f1 prints: percent to one decimal, as a float (None stays None)."""
    return None if v is None else round(v * 100, 1)


def flagged(d):  # same truthiness as plot_linear_probes.flagged: the flag is a string such as "train+val+test", or null
    return bool((d.get("flags") or {}).get("rad_dino_pretraining_overlap"))


def numish(s):
    for conv in (int, float):
        try:
            return conv(s)
        except ValueError:
            pass
    return s


def seed_parts(rc):
    """["seed N"] for a seed other than the default 100, else []: a label or a title names the seed only when it is not the usual one."""
    return [] if numish(rc["seed"]) == 100 else [f"seed {rc['seed']}"]


def seen(path):
    st = Path(path).stat()
    MANIFEST["inputs"][str(path)] = {"mtime": datetime.fromtimestamp(st.st_mtime).astimezone().isoformat(timespec="seconds"), "size": st.st_size}


def rd(path, delimiter=","):
    seen(path)
    with open(path, newline="", encoding="utf-8") as f:  # the evaluation, ten-crop and SSL CSVs are CRLF: newline="" is the csv-module way
        return list(csv.DictReader(f, delimiter=delimiter))


def rj(path):
    seen(path)
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def wrote(path):
    MANIFEST["outputs"].append(str(path))


# ---------------------------------------------------------------- loaders

def run_table(root, container):
    runs = [{"id": i, "dir": d, "label": l, "group": g, "recipe": None} for i, d, l, g in FIXED_RUNS]
    sweep, palette = [], dict(ARM_COLORS)
    for r in rd(container / "sweep" / "wave1.tsv", "\t"):  # read in the table's own order (an id need not be numeric, S2b); that order only colours the extra arms
        w = r["weights"]
        r["seed"] = r.get("seed") or "100"  # a table without the seed column (or an empty cell) is a seed-100 table
        r["init_tag"] = "stock" if "dinov3_vitb16_pretrain" in w else "DINO-X Cosine e30" if w == "STAGEA_FINAL" else "DINO-X ConstLR e30" if "constlr" in w else w
        parts = [f"{r['opt']} lr {r['lr']}", f"wd {r['wd']}", f"ld {r['layer_decay']}", f"pe {r['patch_embed_mult']}", f"init {r['init_tag']}", *seed_parts(r)]
        label = arm_label(r["id"]) or f"{r['id']} ({', '.join(parts)})"  # an arm with a fixed label (run_labels.ARM_LABELS) has it everywhere
        run = {"id": r["id"], "dir": f"all4_cyclic_dinov3_sweep_{r['id']}_seed_{r['seed']}", "label": label, "group": "sweep", "recipe": r}
        if r["id"] not in palette:  # an arm the palette lacks: the next of EXTRA_COLORS (see ARM_COLORS)
            palette[r["id"]] = run["color"] = EXTRA_COLORS[(len(palette) - len(ARM_COLORS)) % len(EXTRA_COLORS)]
        sweep.append(run)
    pos = {i: n for n, i in enumerate(with_extras(SWEEP, {r["id"]: r for r in sweep}))}  # the presentation order, the deck's: SWEEP with the extra adamw arms after S16
    runs += sorted(sweep, key=lambda r: pos.get(r["id"], len(pos)))  # stable: an extra that is not adamw goes last, in the table's order
    live = []
    for r in runs:
        if (root / r["dir"]).is_dir():
            live.append(r)
        else:
            warn(f"{r['id']}: run dir missing, skipped ({root / r['dir']})")
    for r in live:
        if r["group"] == "sweep" and r["id"] not in SWEEP and r["recipe"]["opt"] != "adamw":
            warn(f"{r['id']}: opt {r['recipe']['opt']} is not adamw and the fixed sweep order does not name the id: left out of the deck, the ledger tables, the comparisons and the sweep-wide figures (its sheets stay)")
    return live


def with_extras(order, by_id):
    """order with the sweep table's extra adamw arms (ids SWEEP does not name, in the table's order, which by_id keeps) inserted after S16, before the SGD arms;
    the known ids keep their fixed order."""
    extra = [i for i, r in by_id.items() if r["group"] == "sweep" and i not in SWEEP and r["recipe"]["opt"] == "adamw"]
    k = order.index("S16") + 1
    return (*order[:k], *extra, *order[k:])


def compare_groups(by_id):
    """COMPARE with the extra adamw arms added to the AdamW group."""
    title, ids = COMPARE["adamw"]
    return {**COMPARE, "adamw": (title, with_extras(ids, by_id))}


def arm_colors(by_id):
    """{arm id: colour}: ARM_COLORS plus the colour run_table gave each extra sweep arm."""
    return {**ARM_COLORS, **{i: r["color"] for i, r in by_id.items() if "color" in r}}


def load_eval(root, run):
    """{(split, dataset): (path, rows)} for every evaluation CSV that exists (a live run may have only some); None when none exists."""
    paths = {(sp, ds): root / run["dir"] / "evaluation" / ds / fn for sp, fn in SPLIT_FILES for ds in DATASETS}
    files = {}
    for (sp, ds), p in paths.items():
        if not p.is_file():
            continue
        rows = []
        for r in rd(p):
            t, s = num(r["teacher"]), num(r["student"])
            if (t is None or s is None) and (sp == "test" or r["metric"] in ("mAUC", "validation_loss")):
                BAD[str(p)] += 1  # per-class AUCs of the validation files may be nan and are never used: not flagged
            rows.append({"cycle": int(r["cycle"]), "epoch": int(r["epoch"]), "point": r["evaluation_point"], "metric": r["metric"], "student": s, "teacher": t})
        files[(sp, ds)] = (p, rows)
    if not files:
        warn(f"{run['id']}: no evaluation yet ({run['dir']})")
        return None
    for key, p in paths.items():
        if key not in files:
            warn(f"{run['id']}: evaluation file missing, run kept with the files that exist: {p}")
    return files


def cycle_table(files):
    """{cycle: {"epoch", "end": {ds: (teacher, student)}, "foc": {...}}} from the test mAUC rows.
    end = the after_CheXpert row (end of the cycle); focused = the after_<own dataset> row; cycle 0's init row is both, for every dataset."""
    tab = {}
    for ds in DATASETS:
        for r in files.get(("test", ds), (None, []))[1]:
            if r["metric"] != "mAUC":
                continue
            t = tab.setdefault(r["cycle"], {"epoch": r["epoch"], "end": {}, "foc": {}})
            if r["point"] in (END, INIT):
                t["end"][ds] = (r["teacher"], r["student"])
            if r["point"] in ("after_" + ds, INIT):
                t["foc"][ds] = (r["teacher"], r["student"])
    return tab


def mean_end(t, k):
    v = [t["end"].get(d, (None, None))[k] for d in DATASETS]
    return None if None in v else sum(v) / len(v)


def pick(tab):
    """(complete cycles, last complete cycle, best cycle by mean teacher). Best compares the sheet's 4-decimal values; the strict > keeps the earliest cycle on ties."""
    done = [c for c in sorted(tab) if mean_end(tab[c], 0) is not None and mean_end(tab[c], 1) is not None]
    best = None
    for c in done:
        if best is None or pct(mean_end(tab[c], 0)) > pct(mean_end(tab[best], 0)):
            best = c
    return done, (done[-1] if done else None), best


def best_of(tab, d):
    """(value, cycle) of the best end-of-cycle teacher value of dataset d; the earliest cycle on ties."""
    best = (None, None)
    for c in sorted(tab):
        v = tab[c]["end"].get(d, (None, None))[0]
        if v is not None and (best[0] is None or pct(v) > pct(best[0])):
            best = (v, c)
    return best


def constlr_tags(probes):
    """[(epoch, tag)] of the constant-LR run's probed exports present in probes, by epoch."""
    return sorted((int(m.group(1)), t) for t in probes if (m := CONSTLR.fullmatch(t)))


def tag_label(tag):
    """A probe tag's label: its PROBE_TAGS entry; else, for a constant-LR export the table does not list, the run's name with its epoch; else the tag itself."""
    if label := dict(PROBE_TAGS).get(tag):
        return label
    return f"{SSL_NAMES['constlr']} e{int(m.group(1))}" if (m := CONSTLR.fullmatch(tag)) else tag


def probe_order(probes):
    """The tags present in probes in display order: the PROBE_TAGS rows other than the constant-LR run's, then that run's exports by epoch."""
    return [t for t, _ in PROBE_TAGS if t in probes and not CONSTLR.fullmatch(t)] + [t for _, t in constlr_tags(probes)]


def traj(probes):
    """The probe-trajectory series [(name, ((tag, epoch), ...))]: the cosine run's three exports and the constant-LR run's exports present, by epoch."""
    return ((SSL_NAMES["main"], (("stageA_e0010", 10), ("stageA_e0020", 20), ("stageA_final", 30))),
            (SSL_NAMES["constlr"], tuple((t, n) for n, t in constlr_tags(probes))))


def ckpt_rows(probes):
    """(tag, label) rows of the checkpoint table: the stock weights, the cosine run's exports, the constant-LR run's exports present (by epoch), the RAD-DINO reference, our RayDINO probe."""
    return [("stock_dinov3", "DINOv3 ViT-B/16 stock"), ("stageA_e0010", f"{SSL_NAMES['main']} e10"), ("stageA_e0020", f"{SSL_NAMES['main']} e20"),
            ("stageA_final", f"{SSL_NAMES['main']} e30"), *[(t, tag_label(t)) for _, t in constlr_tags(probes)], ("rad_dino", "RAD-DINO (reference)"), ("raydino", "RayDINO (our probe)")]


def load_probes(root):
    """{tag: {dataset: json}} for the tags of PROBE_TAGS and every constant-LR export on disk (results/constlr_e<NNNN>/); smoke subsets (max_images set) are skipped with a note."""
    probes, present, base = {}, set(), root / "results"
    for tdir in sorted(p for p in base.iterdir() if p.is_dir()) if base.is_dir() else []:
        docs = {ds: rj(tdir / f"{ds}.json") for ds in DATASETS if (tdir / f"{ds}.json").is_file()}
        smoke = [d["max_images"] for d in docs.values() if d.get("max_images") is not None]
        present.add(tdir.name)
        if smoke:
            warn(f"probe tag {tdir.name}: smoke subset (max_images={smoke[0]}), skipped")
        elif tdir.name in dict(PROBE_TAGS) or CONSTLR.fullmatch(tdir.name):
            probes[tdir.name] = docs
        else:
            warn(f"probe tag {tdir.name}: not in the tag table, skipped")
    for tag in [*dict(PROBE_TAGS), *(t for _, t in constlr_tags(probes) if t not in dict(PROBE_TAGS))]:
        if tag in probes:
            for ds in DATASETS:
                if ds not in probes[tag]:
                    warn(f"probe tag {tag}: {ds}.json missing")
        elif tag not in present:
            warn(f"probe tag {tag}: no results under {base}")
    return probes


def offline_files(root, runs):
    """[(run id, cycle, protocol, path)]: the table runs' offline evaluation files, cycle_<NNNN>_<run dir>.csv (protocol tencrop) and
    cycle_<NNNN>_<run dir>.centercrop.csv (protocol centercrop). A file counts whether or not it holds rows."""
    by_dir = {r["dir"]: r["id"] for r in runs}  # only runs of the table: the Step VI and Practice 1 files are ignored
    found = []
    for p in root.glob("cycle_*.csv"):
        m = re.fullmatch(r"cycle_(\d+)_(.+?)(\.centercrop)?\.csv", p.name)
        if m and m.group(2) in by_dir:
            found.append((by_dir[m.group(2)], int(m.group(1)), "centercrop" if m.group(3) else "tencrop", p))
    return found


def tencrop(root, runs):
    """The offline evaluations of the saved weights: [run, cycle, dataset, model, metric, class, percent, protocol] from the offline_files of the table runs.
    A center-crop file without a row (empty, or only its header) is warned about: tencrop_pairs then leaves that run and cycle's center cells empty."""
    rank = {r["id"]: i for i, r in enumerate(runs)}
    rows = []
    for rid, cycle, protocol, p in offline_files(root, runs):
        got = rd(p)
        if not got and protocol == "centercrop":
            warn(f"{p}: no rows; the center-crop cells of that run and cycle are left empty")
        for r in got:
            v = num(r["value"])
            if v is None:
                BAD[str(p)] += 1
            rows.append([rid, cycle, r["dataset"], r["model"], r["metric"], r["class"] or None, pct(v), protocol])
    rows.sort(key=lambda x: (rank[x[0]], x[1], PROTOCOLS.index(x[7])))
    return rows


def tencrop_pairs(ten, tabs, by_id, files=()):
    """Rows of the tencrop_vs_center sheet: one per (run, cycle, dataset) with an offline ten-crop mAUC, each model's ten-crop value next to its
    center-crop value (percent; delta = ten-crop minus center). The center value is the offline center-crop evaluation of the same weights when
    the run and cycle have an offline center-crop file (center_source "offline"), else the in-training end-of-cycle value (center_source
    "in-training"): the source is chosen per run and cycle, so a dataset that file lacks is left empty, never filled from the training evaluation.
    A model with a ten-crop value but no readable center-crop value (entry missing, or its value unreadable) gets empty center and delta cells
    and one warning for the row, naming run, cycle and dataset; the other model keeps its values. A center-crop file of offline_files (files) with no
    row at all is still the source of its run and cycle: every dataset's center cells stay empty, and tencrop() has warned about the file once."""
    got = {tuple(r[:4]): r[6] for r in ten if r[4] == "mAUC" and r[7] == "tencrop"}
    off = {tuple(r[:4]): r[6] for r in ten if r[4] == "mAUC" and r[7] == "centercrop"}
    offline_cycles = {(r[0], r[1]) for r in ten if r[7] == "centercrop"}  # runs and cycles with an offline center-crop file: any row of it
    empty = {(rid, cyc) for rid, cyc, protocol, _ in files if protocol == "centercrop"} - offline_cycles  # ...or the file alone, when it has no row
    offline_cycles |= empty
    rank, dsrank = {i: n for n, i in enumerate(by_id)}, {d: n for n, d in enumerate(DATASETS)}
    rows = []
    for rid, cyc, ds in sorted({k[:3] for k in got}, key=lambda k: (rank[k[0]], k[1], dsrank.get(k[2], len(DATASETS)))):  # a dataset outside DATASETS sorts last
        offline = (rid, cyc) in offline_cycles  # the same weights were also center-cropped offline: this file is the source for every dataset of the cycle
        end = tabs.get(rid, {}).get(cyc, {}).get("end", {}).get(ds, (None, None))  # in-training (teacher, student) at after_CheXpert, fractions; None = unreadable or absent
        cells, lacking = [], []
        for k, model in enumerate(("teacher", "student")):
            c, t = (off.get((rid, cyc, ds, model)) if offline else pct(end[k])), got.get((rid, cyc, ds, model))
            if c is None and (rid, cyc, ds, model) in got:
                lacking.append(model)
            cells += [c, t, None if c is None or t is None else round(t - c, 2)]
        if lacking and (rid, cyc) not in empty:  # a file with no row has its own warning
            warn(f"{rid}: ten-crop cycle {cyc} {ds}: no {'offline center-crop' if offline else 'center-crop end-of-cycle'} value for {' and '.join(lacking)}; center and delta left empty")
        rows.append([rid, by_id[rid]["label"], cyc, ds, *cells, "offline" if offline else "in-training"])
    return rows


def ssl_epochs(root):
    out = []
    for run in SSL_RUNS:
        p = root / f"ssl_dinov3_cxr_{run}_seed_100" / "ssl_losses.csv"
        if not p.is_file():
            warn(f"ssl run {run}: {p} missing")
            continue
        by = {}
        for r in sorted(rd(p), key=lambda r: int(r["step"])):
            by.setdefault(int(r["epoch"]), []).append(r)

        def mean(g, k):
            v = [num(x[k]) for x in g]
            BAD[str(p)] += v.count(None)
            v = [x for x in v if x is not None]
            return statistics.fmean(v) if v else None

        def last(g, k):
            v = num(g[-1][k])
            if v is None:
                BAD[str(p)] += 1
            return v

        for e, g in sorted(by.items()):
            med = statistics.median([len(x) for e2, x in by.items() if e2 != e] or [len(g)])
            if len(g) < med / 2:  # a run still logging (or stopped) mid-epoch leaves a short trailing epoch: not an epoch mean
                warn(f"{p}: epoch {e} dropped as incomplete: {len(g)} logged steps, the run's other epochs have a median of {med:g}")
                continue
            out.append({"run": run, "epoch": e, "steps": len(g), **{k: mean(g, k) for k in LOSS_COLS},
                        "lr_end": last(g, "lr"), "teacher_temp_end": last(g, "teacher_temp"), "step_time_mean": mean(g, "step_time")})
    return out


# ---------------------------------------------------------------- workbook

def order(cycle, point, ds):
    return cycle, POINTS.index(point) if point in POINTS else len(POINTS), DATASETS.index(ds)


def status_of(rc, files, last):
    if rc is None:
        return None
    if files is None:
        return "no evaluation yet"
    return "complete" if len(files) == N_FILES and last == int(rc["cycles"]) else "in progress"


def input_size_of(root, run):
    """The run's input size in px: input_size of its lowest-cycle models/weights/epoch_*/manifest.json in the mirror (224 for a manifest from before the key);
    without a mirrored manifest, the sweep table's input_size (what the launcher passed); else 224, what every run trained at before the column."""
    found = sorted((root / run["dir"] / "models" / "weights").glob("epoch_*/manifest.json"))
    if found:
        try:
            return rj(found[0]).get("input_size", 224)
        except (OSError, ValueError, AttributeError) as e:
            warn(f"{run['id']}: unreadable {found[0]} ({e}): input_size left empty")
            return None
    rc = run["recipe"]
    return numish(rc["input_size"]) if rc and rc.get("input_size") else 224


def cycle0_kind(root, run, files, ten):
    """The run's cycle-0 point: None when it has none (no cycle-0 rows in its evaluation CSVs, no cycle_0000 offline file); "reconstructed init" when
    models/weights/epoch_0000/reconstructed_from_seed.json exists (the point is then in the offline cycle_0000 files; the file is an input of the build, so
    build_manifest.json lists it whenever it exists); else "init" (the trainer evaluated it)."""
    marker = root / run["dir"] / "models" / "weights" / "epoch_0000" / "reconstructed_from_seed.json"
    reconstructed = marker.is_file()
    if reconstructed:
        seen(marker)
    if not (any(r["cycle"] == 0 for _, rows in (files or {}).values() for r in rows) or any(r[0] == run["id"] and r[1] == 0 for r in ten)):
        return None
    return "reconstructed init" if reconstructed else "init"


def end_values(t, conv=pct):
    return [conv(t["end"].get(d, (None, None))[k]) for k in (0, 1) for d in DATASETS]


def run_sheet(tab):
    rows = []
    for c in sorted(tab):
        t = tab[c]
        foc = [pct(t["foc"].get(d, (None, None))[k]) for k in (0, 1) for d in DATASETS]
        rows.append([c, t["epoch"], *end_values(t), *foc, pct(mean_end(t, 0)), pct(mean_end(t, 1))])
    return ["cycle", "epoch", *END_COLS, *FOC_COLS, "T_mean_end", "S_mean_end"], rows


def ledger_sheets(tabs, picks, ids, status):
    """The PDF ledger's three results tables (tables_html) as sheets of numbers: Last evaluated cycle; Best cycle by mean teacher, with the change from the last cycle; Best per dataset (teacher).
    ids = the ledger's runs in its order, status = {run id: text}. Values are p1 numbers, the ones the ledger prints; a run with no complete cycle has no row in the first two."""
    one, vals = "0.0", lambda i, c: end_values(tabs[i][c], p1)
    head = ["Run", "Cycle", *[f"{w} {HEAD_DS[d]}" for w in ("Teacher", "Student") for d in DATASETS], "Status"]
    delta = [f"Δ best − last Teacher {HEAD_DS[d]}" for d in DATASETS]
    last = [[i, picks[i][1], *vals(i, picks[i][1]), status[i]] for i in ids if picks[i][1] is not None]
    prev = {r[0]: r[2:6] for r in last}  # the last cycle's teacher cells, for the deltas
    best, per = [], []
    for i in ids:
        if picks[i][2] is not None:
            vs = vals(i, picks[i][2])
            best.append([i, picks[i][2], *vs, status[i], *[None if a is None or b is None else round(a - b, 1) for a, b in zip(vs[:4], prev[i])], picks[i][1]])
        row = [i]
        for d in DATASETS:
            v, c = best_of(tabs[i], d)
            row += [p1(v), c]
        per.append(row)
    bhead = ["Run", *[h for d in DATASETS for h in (f"{HEAD_DS[d]} best", "cycle")]]
    return [("Last evaluated cycle", head, last, dict.fromkeys(head[2:10], one)),
            ("Best cycle by mean teacher", [*head, *delta, "Last cycle"], best, dict.fromkeys([*head[2:10], *delta], one)),
            ("Best per dataset (teacher)", bhead, per, {h: one for h in bhead if h.endswith(" best")})]


def ledger_disagreements(sheets, ids, tabs):
    """Where the three ledger sheets (ledger_sheets) differ from the summary sheet each one re-presents, cell for cell, as readable strings ([] when they agree): the rows (the ledger's runs that
    the summary has, in the ledger's order), the cycles, the values, Status (as the runs sheet gives it), the deltas (best minus last of the two ledger sheets) and Last cycle. A value is checked twice:
    exactly against the one-decimal rounding of the raw fraction in tabs (build's {run id: cycle_table}) at the summary's cycle, and within the rounding of the summary's 4-decimal value. The tolerance alone
    cannot tell the two neighbouring one-decimal values apart when the summary sits at a midpoint (12.3500 is within it of both 12.3 and 12.4); the raw fraction can."""
    sh = {n: (h, r) for n, h, r, _ in sheets}
    runs = {r[0]: dict(zip(sh["runs"][0], r)) for r in sh["runs"][1]}
    tol = 0.05 + 0.00005 + 1e-9  # half a unit of the one-decimal value plus the summary's own 4-decimal rounding; both come from the same fraction (the bound between the two sheets; exact below decides the value)
    bad = []

    def diff(name, i, what, new, old, close=False):
        if not ((new is None) == (old is None) and (new is None or (abs(new - old) <= tol if close else new == old))):
            bad.append(f"{name}: {i} {what}: {new} where {old} was expected")

    def exact(name, i, what, new, c, n):  # new must be the one-decimal rounding of the raw n-th end_values entry (teacher x4, student x4) of run i at cycle c; no cycle, no value; a run or cycle tabs lacks is reported
        t = tabs.get(i, {}).get(c)
        if t is None and c is not None:
            bad.append(f"{name}: {i} {what}: no raw values for cycle {c} to check against")
        else:
            diff(name, i, f"{what} (raw fraction, cycle {c})", new, None if t is None else end_values(t, p1)[n])

    def pairs(name, summary):  # the ledger sheet's rows with the summary's, by run id; the row set and its order are checked here
        new, old = sh[name][1], {r[0]: r for r in sh[summary][1]}
        want = [i for i in ids if i in old]
        if [r[0] for r in new] != want:
            bad.append(f"{name}: runs {[r[0] for r in new]} where {summary} gives {want}")
        return [(r, old[r[0]]) for r in new if r[0] in old]

    last_cycle, prev = {r[0]: r[2] for r in sh["summary_last"][1]}, {r[0]: r for r in sh["Last evaluated cycle"][1]}
    for name, summary, deltas in (("Last evaluated cycle", "summary_last", False), ("Best cycle by mean teacher", "summary_best_mean", True)):
        head = sh[name][0]
        for r, o in pairs(name, summary):  # r = Run, Cycle, 4 teacher, 4 student, Status[, 4 deltas, Last cycle]; o = id, label, cycle, 4 teacher, 4 student, T_mean_end
            i, run = r[0], runs.get(r[0], {})
            diff(name, i, "Cycle", r[1], o[2])
            for n, (h, a, b) in enumerate(zip(head[2:10], r[2:10], o[3:11])):
                diff(name, i, h, a, b, True)
                exact(name, i, h, a, o[2], n)  # at the summary's cycle, not the ledger's own
            diff(name, i, "Status", r[10], f"in progress (c{run['last_cycle']} of {run['cycles']})" if run.get("status") == "in progress" else run.get("status"))
            if deltas:
                p = prev.get(i, [None] * 11)
                for k in range(4):
                    diff(name, i, head[11 + k], r[11 + k], None if r[2 + k] is None or p[2 + k] is None else round(r[2 + k] - p[2 + k], 1))
                diff(name, i, "Last cycle", r[15], last_cycle.get(i))
    name = "Best per dataset (teacher)"
    for r, o in pairs(name, "summary_best_per_dataset"):  # r = Run, then best and cycle per dataset; o = id, label, then the same
        for k, what in enumerate(f"{HEAD_DS[d]} {w}" for d in DATASETS for w in ("best", "cycle")):
            diff(name, r[0], what, r[1 + k], o[2 + k], k % 2 == 0)
            if k % 2 == 0:
                exact(name, r[0], what, r[1 + k], o[3 + k], k // 2)  # a best value: the teacher fraction of that dataset at the summary's cycle for it
    return bad


def eval_rows(run, files, split, metric_ok, with_class):
    rows = []
    for ds in DATASETS:
        for r in files.get((split, ds), (None, []))[1]:
            if not metric_ok(r["metric"]):
                continue
            cls = [r["metric"][len("AUC_"):]] if with_class else []
            rows.append([run["id"], run["label"], split, r["cycle"], r["epoch"], r["point"], ds, *cls,
                         r["point"] in ("after_" + ds, INIT), r["point"] in (END, INIT), pct(r["student"]), pct(r["teacher"])])
    rows.sort(key=lambda x: order(x[3], x[5], x[6]))
    return rows


def probe_sheets(probes):
    tags = probe_order(probes)
    classes = []
    for tag in tags:
        for ds in DATASETS:
            classes += [c for c in probes[tag].get(ds, {}).get("classes", []) if c not in classes]
    head = ["tag", "label", "dataset", "n_train", "n_test", "runs", "test_mAUC_mean", "test_mAUC_std", "overlap_flag", "pretraining_overlap_train", "warnings",
            *[f"{c}_{s}" for c in classes for s in ("mean", "std")]]
    rows, per_run = [], []
    for tag in tags:
        for ds, d in probes[tag].items():
            rows.append([tag, tag_label(tag), ds, d["n"]["train"], d["n"]["test"], d["runs"], pct(d["test_mAUC_mean"]), pct(d["test_mAUC_std"]), flagged(d), d.get("pretraining_overlap_train"),
                         "; ".join(map(str, d.get("warnings") or [])),
                         *[pct(d[k].get(c)) for c in classes for k in ("test_AUC_mean", "test_AUC_std")]])
            per_run += [[tag, ds, p["run"], p["seed"], p["best_epoch"], pct(p["test_mAUC"])] for p in d["per_run"]]
    return head, rows, per_run


def write_workbook(path, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for name, head, rows, fmt in sheets:
        ws = wb.create_sheet(name)
        ws.append(head)
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{ws.max_row}"
        for j, h in enumerate(head, 1):
            if h in fmt:
                for (cell,) in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                    cell.number_format = fmt[h]
            width = max([len(str(h))] + [len(str(r[j - 1])) for r in rows[:200] if r[j - 1] is not None])
            ws.column_dimensions[get_column_letter(j)].width = min(width + 2, 70)
    wb.save(path)


def readme_lines(root, ids, now):
    return [
        f"DINO-X — numbers workbook ({now:%Y-%m-%d}). Every cell is read from the files under the mirror root; an unreadable value is left empty.",
        "Project: DINO-X. Earlier files were titled Step VIII (Demo 1); same data.",
        f"Mirror root: {root}",
        f"Built: {now:%Y-%m-%d %H:%M:%S %Z} (local time)",
        f"Script: {Path(__file__).resolve()}",
        "Scale: AUC and mAUC are stored x100 (percent) rounded to 4 decimals; validation_loss and the SSL losses are stored as logged. T_ = teacher, S_ = student.",
        "Training order inside a cycle is VinDrCXR, ChestXray14, MIMIC, CheXpert. evaluation_point after_<D> is the evaluation right after dataset D's pass; after_CheXpert is the end of the cycle.",
        "Dataset names: " + "; ".join(f"{d} = {TITLE_DS[d]} (column id {SHORT[d]})" for d in DATASETS)
        + ". The dataset columns hold the directory ids, the workbook headers (T_CXR14_end) the column ids; the deck, figures and HTML tables use the display names.",
        "Focused value for dataset D (CSE 494 IP deck p. 29) = the row at after_D. Unfocused = the rows at the other three evaluation points.",
        "End of cycle = the after_CheXpert rows. They are the ledger's headline numbers: focused for CheXpert, unfocused for the other three datasets.",
        "Last cycle = the latest cycle whose four after_CheXpert test mAUC rows exist for both models. Ties for the best go to the earliest cycle.",
        "Graphs start at cycle 2: the trainer evaluates every 2 cycles (eval_every 2 for these arms, 1 for the input-size arms S18 and S19, whose graphs start at cycle 1; sheet runs lists each arm's eval_every), so cycle 1, the first warmup cycle, has no evaluation point; cycle 2, the second warmup cycle, is evaluated and drawn.",
        "The x-axis of every curve starts at the first evaluated point: cycle 2 for the sweep arms (cycle 1 has no evaluation), cycle 0 for a run with cycle-0 rows.",
        "Comparison and sweep-wide figures draw the teacher only; per-run figures draw student and teacher.",
        "An adamw arm of sweep/wave1.tsv whose id the fixed arm order does not name (a second-seed arm such as S2b, or the input-size arms S18 and S19) follows S16 in the deck, the ledger tables and the sweep-wide figures and joins the AdamW comparison; "
        "it takes the next of tab20 entries 16, 17, 18, 19 and 15 (the grays last) in the table's file order, which does nothing else; a title names the seed only when it is not 100.",
        "Cycle 0 is the initial model before any training: evaluation_point init, epoch 0, with a student value and a teacher value of its own (two separate model builds, so their random heads "
        "differ). Its rows are both the focused and the end-of-cycle rows of every dataset (focused and end_of_cycle are TRUE), it counts as a cycle in the per-run sheets, the summaries and the "
        "best-cycle pick, and the graphs draw it at x = 0; it is told from cycle 1 (also epoch 0) by its cycle number and evaluation_point, never by epoch. runs.cycle0 says init when the trainer "
        "evaluated it, or reconstructed init when models/weights/epoch_0000/reconstructed_from_seed.json exists for the run (the cycle-0 weights were rebuilt from the seed and evaluated offline: "
        "cycle_0000_<run dir>.csv and cycle_0000_<run dir>.centercrop.csv).",
        "Two offline protocols evaluate the saved weights of every cycle: ten-crop (Resize 256 → TenCrop 224, sigmoid averaged over the 10 crops) and center crop (Resize 256 → CenterCrop 224, "
        "the protocol of the in-training evaluation, which runs every second cycle and, where an init point exists, at cycle 0). Both are in sheet offline_eval, told apart by protocol. "
        "A run trained at another input size is evaluated at its own (runs.input_size; resize = round(input_size * 8 / 7)): S18 Resize 512 → TenCrop / CenterCrop 448 px, S19 Resize 585 → 512 px.",
        "Sweep-wide center-crop figures draw the in-training end-of-cycle value where there is one and the offline center-crop value where there is none (cycle 0 when it was not evaluated in "
        "training, and odd cycles): a solid marker is in-training, a hollow marker of the same colour is offline; the legend note reads: hollow = offline evaluation of the saved weights.",
        "Sheets:",
        "runs: one row per run; recipe columns from sweep/wave1.tsv for sweep arms (init as in the TSV, init_tag = the tag derived from weights); status = complete, in progress (last cycle below the target, or evaluation CSVs missing) or no evaluation yet (no CSV at all); cycle0 = init (or reconstructed init) when the run's evaluation CSVs hold cycle-0 rows, empty otherwise; input_size = the crop in px, from the run's manifest.json in the mirror (224 when the manifest predates the key), else from the table's input_size, else 224; seed, eval_every and eval_cycle0 (the last three columns) are the table's own columns "
        "(a table without them is a seed-100 table, the other two empty), and the sweep arms' run directory is all4_cyclic_dinov3_sweep_<id>_seed_<seed>; the sweep arms follow the deck's order (the fixed arm order, extra arms after S16), not the table's.",
        "Last evaluated cycle: the PDF ledger's table of that name, as numbers: the ledger's runs in its order (the early arms, then the sweep arms; probe_a and probe_b, which the ledger tables leave out, are in the summary sheets only), the end-of-cycle test mAUC at the last cycle, teacher then student, in percent to one decimal (the number the ledger prints); Status = complete, or in progress (cN of M): last cycle N of the M cycles of sweep/wave1.tsv, empty for the early arms. The same numbers as summary_last; the build checks it.",
        "Best cycle by mean teacher: the ledger's Best cycle by mean teacher mAUC table: the same columns at summary_best_mean's cycle, then Δ best − last per dataset (teacher, this sheet's value minus Last evaluated cycle's, percentage points) and the Last cycle. The same numbers as summary_best_mean; the build checks it.",
        "Best per dataset (teacher): the ledger's Best teacher value per dataset table: per run and dataset the best end-of-cycle teacher value, percent to one decimal, and the cycle it occurred. The same numbers as summary_best_per_dataset; the build checks it.",
        "cyclic_mAUC: every mAUC row, test and validation: one row per run, split, cycle, evaluation point and dataset.",
        "cyclic_classAUC: every per-class AUC row of the test split.",
        "validation_loss: the validation_loss rows of the validation files.",
        f"{', '.join(ids)}: one sheet per run (test split), one row per cycle: end-of-cycle values, focused values, and the mean of the four end-of-cycle values.",
        "summary_last: the end-of-cycle values at the last cycle of each run.",
        "summary_best_mean: the end-of-cycle values at the cycle with the highest mean teacher value.",
        "summary_best_per_dataset: the best end-of-cycle teacher value per dataset and the cycle it occurred.",
        "probes: frozen linear probes per checkpoint and dataset: mean and s.d. over the probe fits, and per class; overlap_flag = RAD-DINO pretraining overlap.",
        "RayDINO's pretraining corpus is PadChest plus the training splits of CheXpert, MIMIC-CXR and ChestX-ray14 (its Table 2); VinDr-CXR is not in it; no test-set overlap, so no dagger; the train-split overlap is recorded in pretraining_overlap_train.",
        "probes_per_run: the individual probe fits.",
        "offline_eval (renamed from tencrop): the offline evaluations of the saved weights at the mirror root, one file per run and cycle: protocol tencrop = cycle_<NNNN>_<run dir>.csv "
        "(eval_d1_6.sh --test_augment true), protocol centercrop = cycle_<NNNN>_<run dir>.centercrop.csv.",
        "tencrop_vs_center: each offline ten-crop mAUC next to the center-crop mAUC of the same run, cycle, dataset and model. center_source = offline when that run and cycle has an "
        "offline center-crop file for the same weights (it is the center value; a dataset the file lacks, or a file with no row, is left empty, with a warning), in-training otherwise (the end-of-cycle value of the training evaluation; center crop = Resize 256 → "
        "CenterCrop 224, every training launcher sets ARK_TEST_AUGMENT=false). delta = ten-crop − center, percentage points.",
        "ssl_losses_epoch: per-epoch means of the DINO-X pretraining (internal tag Stage A) loss logs; run = the pretraining run directory ssl_dinov3_cxr_<run>_seed_100 (main = DINO-X Cosine (cosine schedule); constlr = DINO-X ConstLR (constant LR); pilot_a and pilot_b = the two pilot runs); an epoch with fewer than half the median step count of the run's other epochs is dropped as incomplete.",
    ]


# ---------------------------------------------------------------- HTML fragment

def td(cls, text):
    return f'<td class="{cls}">{text}</td>' if cls else f"<td>{text}</td>"


def th(cls, text):
    return f'<th class="{cls}">{html.escape(text)}</th>' if cls else f"<th>{html.escape(text)}</th>"


def html_table(tid, minw, head, body):
    return (f'<div class="scroll">\n<table id="{tid}" style="min-width: {minw}px">\n<thead>\n<tr>' + "".join(th(c, t) for t, c in head)
            + "</tr>\n</thead>\n<tbody>\n" + "\n".join(body) + "\n</tbody>\n</table>\n</div>")


def run_rows(tabs, cycles, sweep=SWEEP):
    body = []
    for title, ids in (("Early arms", EARLY), ("Sweep wave 1", sweep)):
        body.append(f'<tr class="group"><td colspan="10">{title}</td></tr>')
        for rid in ids:
            c = cycles.get(rid)
            if c is None:
                continue
            t = tabs[rid][c]
            cells = [td("r" if k == 0 else "r s" + (" sep" if i == 0 else ""), f1(t["end"].get(d, (None, None))[k])) for k in (0, 1) for i, d in enumerate(DATASETS)]
            body.append(f'<tr>{td("id", rid)}{td("cyc", f"c{c}")}{"".join(cells)}</tr>')
    return body


def probe_cell(d, std=False):
    if d is None:
        return ""
    return f"{100 * d['test_mAUC_mean']:.{2 if std else 1}f}" + (f" ({100 * d['test_mAUC_std']:.2f})" if std else "") + ("†" if flagged(d) else "")


def tables_html(tabs, picks, probes, pairs, sweep=SWEEP):
    head_run = [("Run", ""), ("Cycle", "")] + [(f"{w} {HEAD_DS[d]}", "r" + (" sep" if (w, d) == ("Student", DATASETS[0]) else "")) for w in ("Teacher", "Student") for d in DATASETS]
    last = {i: p[1] for i, p in picks.items()}
    best = {i: p[2] for i, p in picks.items()}
    best_ds = []
    for rid in (*EARLY, *sweep):
        if rid in tabs:
            cells = []
            for d in DATASETS:
                v, c = best_of(tabs[rid], d)
                cells.append(td("r", "" if v is None else f"{f1(v)} <small>c{c}</small>"))
            best_ds.append(f'<tr>{td("id", rid)}{"".join(cells)}</tr>')
    ckpt = []
    for tag, label in ckpt_rows(probes):
        docs = probes.get(tag)
        if docs is None:
            continue
        four = [docs.get(d) for d in DATASETS]
        mean = "" if None in four else f"{sum(100 * d['test_mAUC_mean'] for d in four) / len(four):.1f}"
        ckpt.append(f'<tr>{td("id", html.escape(label))}{"".join(td("r", probe_cell(d)) for d in four)}{td("r", mean)}</tr>')
    frozen = []
    for tag in probe_order(probes):
        frozen.append(f'<tr>{td("id", html.escape(tag_label(tag)))}{"".join(td("r", probe_cell(probes[tag].get(d), std=True)) for d in DATASETS)}</tr>')
    ds_head = [(HEAD_DS[d], "r") for d in DATASETS]
    by_pair = {}  # (run, cycle) -> {dataset: (center, ten-crop)}, teacher; the pair rows are in sheet order
    for r in pairs:
        if r[0] in (*EARLY, *sweep):  # the runs the other ledger tables list: an arm outside them (an extra that is not adamw) keeps its sheet rows only
            by_pair.setdefault((r[0], r[2]), {})[r[3]] = (r[4], r[5])
    head_tc = [("Run", ""), ("Cycle", "")] + [(f"{HEAD_DS[d]} {w}", "r sep" if (d, w) == (DATASETS[0], "center") else "r") for d in DATASETS for w in ("center", "ten")]
    tc = []
    for (rid, c), by_ds in by_pair.items():
        vals = [v for d in DATASETS for v in by_ds.get(d, (None, None))]
        cells = [td("r sep" if i == 0 else "r", f1(None if v is None else v / 100)) for i, v in enumerate(vals)]  # the pair rows are percent, f1 takes a fraction
        tc.append(f'<tr>{td("id", rid)}{td("cyc", f"c{c}")}{"".join(cells)}</tr>')
    parts = [
        ("Last evaluated cycle", html_table("ft-last", 760, head_run, run_rows(tabs, last, sweep))),
        ("Best cycle by mean teacher mAUC", html_table("ft-bestmean", 760, head_run, run_rows(tabs, best, sweep))),
        ("Best teacher value per dataset, with the cycle it occurred", html_table("ft-best", 640, [("Run", ""), *ds_head], best_ds)),
        ("Linear-probe results by checkpoint", html_table("ssl-probes", 640, [("Checkpoint", ""), *ds_head, ("Mean", "r")], ckpt)),
        ("Our frozen probes (test mAUC, percent, probe s.d. over 10 fits in parentheses)",
         html_table("frozen-probes", 760, [("Backbone", ""), *ds_head[:-1], ("CheXpert (5)", "r")], frozen)),
        ("Ten-crop vs center crop (test mAUC, percent, teacher; end of cycle)", html_table("ft-tencrop", 760, head_tc, tc)),
    ]
    return "\n\n".join(f"<h3>{html.escape(t)}</h3>\n{tbl}" for t, tbl in parts) + "\n"


# ---------------------------------------------------------------- figures and deck

def pyplot():
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    return plt


def save_fig(plt, fig, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.tight_layout()
    fig.savefig(path, dpi=180, bbox_inches="tight")
    plt.close(fig)
    wrote(path)


def run_script(script, *args):
    r = subprocess.run([sys.executable, "-B", str(PLOT_DIR / script), *map(str, args)], capture_output=True, text=True)
    if r.returncode:
        sys.exit(f"{script} failed (exit {r.returncode}):\n{r.stderr.strip()}")
    for line in r.stderr.splitlines():
        if line.strip():
            warn(f"{script}: {line.strip()}")


def copy_png(src, dst):
    if not src.is_file():
        warn(f"figure missing: {src}")
        return
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(src, dst)
    wrote(dst)


def cyclic_figs(root, runs, fig_dir, tmp):
    # ponytail: one subprocess per run, sequential (about 2 s each); parallelise only if the build time ever matters
    for run in runs:
        run_script("plot_cyclic_run.py", root / run["dir"], "--format", "png", "--dpi", 180, "--output-root", tmp)
        for (kind, split), ds in ((k, d) for k in FIG_KINDS for d in DATASETS):
            copy_png(tmp / run["dir"] / kind / split / f"{ds}.png", fig_dir / "cyclic" / run["id"] / f"{kind}_{split}_{ds}.png")


def compare_figs(root, by_id, have, fig_dir, tmp):
    center = (by_id.get(CENTER) or {}).get("recipe")
    colors = arm_colors(by_id)
    for group, (_, ids) in compare_groups(by_id).items():
        order = [i for i in sorted(ids, key=lambda i: i == CENTER) if i in have]  # the centre arm last: the script draws in argument order
        dirs = [root / by_id[i]["dir"] for i in order]
        if len(dirs) < 2:
            warn(f"comparison {group}: fewer than two runs with evaluations, skipped")
            continue
        out = tmp / f"cmp_{group}"
        run_script("compare_cyclic_runs.py", "--split", "test", "--format", "png", "--dpi", 180, "--models", "teacher", "--full-length", "--output-root", out, *dirs,
                   "--colors", *[f"{by_id[i]['dir']}={colors[i]}" for i in order],  # the arm palette: an arm has the same colour in every figure
                   "--labels", *[f"{by_id[i]['dir']}={short_title(by_id[i], center)}" for i in order])  # the sweep-wide figures' legend strings: the arm id with what sets it apart
        for ds in DATASETS:  # the script names its folder from the run dirs: one folder per fresh output root
            copy_png(next(iter(out.glob(f"*/test/{ds}.png")), out / "missing" / f"{ds}.png"), fig_dir / "compare" / group / f"{ds}.png")


def probe_page_tags(probes):
    """The tags the probe pages draw: the PROBE_PAGE_TAGS that have results, then the latest probed constant-LR export (the highest constlr_e<N> present)."""
    found = constlr_tags(probes)
    return [t for t in PROBE_PAGE_TAGS if t in probes] + ([found[-1][1]] if found else [])


def probe_pages(root, probes, fig_dir, date):
    """plot_linear_probes' pages as PNGs, in page order: [(page name, path)]. Its PdfPages is swapped for a stub that writes PNGs.
    The pages draw one export per SSL run (probe_page_tags); the workbook and the tables keep every epoch."""
    import plot_linear_probes as P

    tags = probe_page_tags(probes)
    if not tags:
        warn("probe pages skipped: none of the page tags has results")
        return []
    out = fig_dir / "probes"
    out.mkdir(parents=True, exist_ok=True)
    pages = []

    @contextlib.contextmanager
    def png_pages(_path):
        def savefig(fig):
            path = out / f"page{len(pages) + 1}.png"
            fig.savefig(path, dpi=180)
            pages.append((fig.get_suptitle().split(":", 1)[0], path))
            wrote(path)

        yield types.SimpleNamespace(savefig=savefig)

    P.PdfPages = png_pages
    argv = [str(root / "results"), "--tags", *[f"{t}={tag_label(t)}" for t in tags], "--datasets", *DATASETS, "--out", str(out / "linear_probes.pdf"),
            "--reference", "=".join(ARK_REF), "--title", f"DINO-X — linear probes ({date})"]
    err = io.StringIO()
    try:
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(err):
            P.main(argv)
    except SystemExit:  # a refusal (smoke subset, cohort mismatch) stops the build with the script's own message
        sys.stderr.write(err.getvalue())
        raise
    for line in err.getvalue().splitlines():
        if line.strip() and "wrote " not in line:
            warn(line.strip())
    return pages


def fig_trajectory(probes, path):
    plt = pyplot()
    from plot_cyclic_run import NATURE_PALETTE

    def val(tag, key):  # (mean, s.d.) in percent for a dataset; for the mean of four the s.d. is in quadrature (assumes independent fits)
        d = probes.get(tag, {})
        if key == "mean":
            if any(ds not in d for ds in DATASETS):
                return None
            return sum(100 * d[ds]["test_mAUC_mean"] for ds in DATASETS) / len(DATASETS), math.sqrt(sum((100 * d[ds]["test_mAUC_std"]) ** 2 for ds in DATASETS)) / len(DATASETS)
        return (100 * d[key]["test_mAUC_mean"], 100 * d[key]["test_mAUC_std"]) if key in d else None

    fig, axes = plt.subplots(2, 3, figsize=(14, 7.5))
    runs = traj(probes)
    for ax, key in zip(axes.flat, (*DATASETS, "mean")):
        at = {}
        for (name, series), color in zip(runs, NATURE_PALETTE):
            pts = [(x, *v) for tag, x in series if (v := val(tag, key))]
            if pts:
                xs, ys, es = zip(*pts)
                ax.errorbar(xs, ys, yerr=es, marker="o", capsize=3, linewidth=1.8, color=color, label=name)
                for x, y in zip(xs, ys):
                    at.setdefault(x, []).append((y, color))
        for x, items in at.items():
            for rank, (y, color) in enumerate(sorted(items, reverse=True)):
                ax.annotate(f"{y:.2f}", (x, y), xytext=(0, 9 if rank == 0 else -9), textcoords="offset points", ha="center",
                            va="bottom" if rank == 0 else "top", fontsize=8, color=color)
        if (stock := val("stock_dinov3", key)):
            ax.axhline(stock[0], color="0.4", linestyle=":", label="DINOv3 ViT-B/16 stock")
        ax.set_title("Mean of four" if key == "mean" else TITLE_DS[key], fontsize=12)
        ax.set_xticks(sorted({x for _, series in runs for _, x in series}))
        ax.margins(x=0.15, y=0.2)
        ax.set_xlabel("Export epoch")
        ax.set_ylabel("Test mAUC (%)")
        ax.grid(axis="y", alpha=0.22, linewidth=0.8)
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)
    handles, labels = axes.flat[0].get_legend_handles_labels()
    axes.flat[5].axis("off")
    axes.flat[5].legend(handles, labels, loc="center", frameon=False, fontsize=12)
    fig.suptitle("DINO-X export: frozen linear-probe test mAUC (error bars: probe s.d.)", fontsize=13)
    save_fig(plt, fig, path)


def fig_loss(ssl, name, title, path):
    plt = pyplot()
    from plot_cyclic_run import NATURE_PALETTE

    key = "lr_end" if name == "lr" else name
    fig, ax = plt.subplots(figsize=(11, 6.2))
    for run, color in zip(("main", "constlr"), NATURE_PALETTE):
        pts = [(r["epoch"], r[key]) for r in ssl if r["run"] == run and r[key] is not None]
        if pts:
            ax.plot(*zip(*pts), color=color, linewidth=2.0, label=SSL_NAMES[run])
    ax.set_title(title, fontsize=13, pad=10)
    ax.set_xlabel("Epoch")
    ax.set_ylabel("lr at the end of the epoch" if name == "lr" else f"{name}: mean over the epoch's logged steps")
    ax.grid(axis="y", alpha=0.22, linewidth=0.8)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.legend(frameon=False)
    save_fig(plt, fig, path)


def short_title(run, center):
    """Slide-title identity of a run: sweep arms are named by what differs from the centre recipe, and by the seed when it is not 100."""
    rc = run["recipe"]
    if rc is None or center is None:
        return run["label"]
    if fixed := arm_label(run["id"]):  # the run label: the same string in titles, legends and the workbook
        return fixed
    if rc is center:
        init = "stock DINOv3" if rc["init_tag"] == "stock" else rc["init_tag"]  # the legend says what the arm starts from
        parts = ["center recipe", f"{init} init", *seed_parts(rc)]
        return f"{run['id']} ({', '.join(parts)})"
    diff = [rc[k] if k == "opt" else f"{RECIPE_ABBR.get(k, k)} {rc[k]}" for k in (*RECIPE_KEYS, "init_tag") if rc[k] != center[k]] + seed_parts(rc)
    return f"{run['id']} ({', '.join(diff)})" if diff else run["id"]


def sweep_series(ten, tabs, picks, ds, ids=SWEEP):
    """(ten-crop, center-crop, hollow) teacher test mAUC against cycle, in ids order (SWEEP unless given): the first two are {run id: [(cycle, percent)]}.
    Ten-crop: the sweep arms that have an offline ten-crop teacher mAUC for ds. Center-crop: every evaluated sweep arm, its in-training end-of-cycle
    teacher value for ds at every cycle that has one (cycle 0's init point included; a cycle still missing other datasets or the student counts, so picks,
    the complete cycles, no longer decides), plus the offline center-crop value at every cycle that has no in-training value (cycle 0 when it was not
    evaluated in training, odd cycles); hollow[run id], for arms that have any, is the set of those offline cycles."""
    tcs, ccs, hollow = {}, {}, {}
    for rid in ids:
        pts = sorted((r[1], r[6]) for r in ten if (r[0], r[2], r[3], r[4], r[7]) == (rid, ds, "teacher", "mAUC", "tencrop") and r[6] is not None)
        if pts:
            tcs[rid] = pts
        cc = [(c, v) for c in sorted(tabs.get(rid, {})) if (v := pct(tabs[rid][c]["end"].get(ds, (None, None))[0])) is not None]
        off = sorted((r[1], r[6]) for r in ten if (r[0], r[2], r[3], r[4], r[7]) == (rid, ds, "teacher", "mAUC", "centercrop") and r[6] is not None and r[1] not in {c for c, _ in cc})
        if cc or off:
            ccs[rid] = sorted(cc + off)
            if off:
                hollow[rid] = {c for c, _ in off}
    return tcs, ccs, hollow


def fig_sweep_all(series, title, labels, colors, path, hollow=None):
    plt = pyplot()
    from matplotlib.lines import Line2D
    from matplotlib.ticker import MaxNLocator
    from plot_cyclic_run import xlim_from_first_point

    hollow = hollow or {}
    fig, ax = plt.subplots(figsize=(11, 6.2))
    for rid, pts in series.items():
        xs, ys = zip(*pts)
        solid = [i for i, x in enumerate(xs) if x not in hollow.get(rid, ())]
        ax.plot(xs, ys, marker="o" if solid else None, markevery=solid or None, markersize=5, linewidth=1.8, color=colors[rid], label=labels[rid], clip_on=False)  # one point: a marker only; unclipped, so the first cycle's markers stay whole on the y-axis
        if len(solid) < len(xs):  # the offline points: hollow markers of the arm's colour, over the line
            ax.plot([x for x in xs if x in hollow[rid]], [y for x, y in pts if x in hollow[rid]], linestyle="none", marker="o", markersize=5,
                    markerfacecolor="white", markeredgecolor=colors[rid], zorder=3, clip_on=False)
    ax.xaxis.set_major_locator(MaxNLocator(integer=True, min_n_ticks=1))  # one integer in view (a one-point or one-cycle series) is enough: whole cycles, never fractional ticks
    xlim_from_first_point(ax)  # the x axis runs from the first to the last drawn cycle, no margin either side (the markers are clip_on=False, so the end markers stay whole)
    ax.set_title(title, fontsize=13, pad=10)
    ax.set_xlabel("Cycle")
    ax.set_ylabel("Test mAUC (%)")
    ax.grid(axis="y", alpha=0.22, linewidth=0.8)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    handles, names = ax.get_legend_handles_labels()
    if any(hollow.get(rid) for rid in series):  # one note, only when hollow markers are drawn
        handles.append(Line2D([], [], marker="o", linestyle="none", markersize=5, markerfacecolor="white", markeredgecolor="0.35"))
        names.append(CC_NOTE)
    ax.legend(handles, names, frameon=False, fontsize=8, bbox_to_anchor=(1.02, 1), loc="upper left")
    save_fig(plt, fig, path)


def sweep_all_figs(ten, tabs, picks, by_id, out):
    """tencrop_<dataset>.png (the sweep arms with offline ten-crop data) and center_<dataset>.png (every evaluated sweep arm; offline center-crop points hollow) in out: teacher test mAUC against cycle."""
    center = (by_id.get(CENTER) or {}).get("recipe")
    ids, colors = with_extras(SWEEP, by_id), arm_colors(by_id)
    labels = {rid: short_title(by_id[rid], center) for rid in ids if rid in by_id}
    for ds in DATASETS:
        tcs, ccs, hollow = sweep_series(ten, tabs, picks, ds, ids)
        for (kind, name), series, hol in zip((("tencrop", "Ten-crop"), ("center", "Center crop")), (tcs, ccs), ({}, hollow)):
            if series:
                fig_sweep_all(series, f"{name} — {TITLE_DS[ds]} — test mAUC, teacher — sweep arms", labels, colors, out / f"{kind}_{ds}.png", hol)
            else:
                warn(f"sweep figure {kind}_{ds}.png skipped: no sweep arm has data for it")


def make_deck(slides, path):
    """slides: [(title, picture)]. A picture is one PNG (it fills the content box) or a list of four PNGs (a 2x2 grid, in list order)."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(40 / 3), Inches(7.5)
    mx, top, bottom, gutter = Inches(0.4), Inches(0.85), Inches(0.3), Inches(0.1)
    aw, ah = prs.slide_width - 2 * mx, prs.slide_height - top - bottom

    def place(s, png, x, y, w, h):  # the picture scaled to fit the w x h box at (x, y), centred in it
        pic = s.shapes.add_picture(str(png), 0, 0)
        k = min(w / pic.width, h / pic.height)
        pic.width, pic.height = int(pic.width * k), int(pic.height * k)
        pic.left, pic.top = int(x + (w - pic.width) / 2), int(y + (h - pic.height) / 2)

    for title, png in slides:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        long = len(title) > 90  # over ~90 characters: 16 pt with wrap on (16 pt holds about 110 characters on one line)
        tf = s.shapes.add_textbox(mx, Inches(0.2), aw, Inches(0.5)).text_frame
        tf.word_wrap = long
        tf.margin_left = tf.margin_right = 0
        tf.text = title
        tf.paragraphs[0].runs[0].font.size = Pt(16 if long else 20)
        if isinstance(png, list):  # 2x2: each picture in its aw/2 x ah/2 cell, a gutter wide between neighbours; a missing PNG leaves its cell empty
            for i, q in enumerate(png):
                if q.is_file():
                    place(s, q, mx + i % 2 * aw / 2 + gutter / 2, top + i // 2 * ah / 2 + gutter / 2, aw / 2 - gutter, ah / 2 - gutter)
                else:
                    warn(f"deck: {title}: picture missing, its cell left empty: {q}")
        else:
            place(s, png, mx, top, aw, ah)
        if len(title) > 110:  # not even 16 pt holds it on one line
            warn(f"deck: title may not fit on one line ({len(title)} chars): {title}")
    prs.save(path)
    wrote(path)
    return len(prs.slides)


def deck_entries(fig_dir, tabs, by_id, pages):
    """[(slide title, PNG or the four per-dataset PNGs)] in deck order. Titles read <subject> — <dataset> — <split> — Teacher; a per-run slide holds all four datasets."""
    center = (by_id.get(CENTER) or {}).get("recipe")
    deck = [(f"Linear probes — {TITLE_DS.get(n, n)}", p) for n, p in pages]
    deck += [("DINO-X probe trajectory", fig_dir / "ssl" / "probe_trajectory.png")]
    deck += [(title, fig_dir / "ssl" / f"losses_{name}.png") for name, title in LOSS_FIGS]
    for rid in with_extras(DECK_ORDER, by_id):
        if rid in tabs:
            for split in ("test", "validation"):
                deck += [(f"{short_title(by_id[rid], center)} — {split} — focused and unfocused — Student and Teacher",
                          [fig_dir / "cyclic" / rid / f"focused_unfocused_{split}_{ds}.png" for ds in DATASETS])]
    for group, (title, _) in COMPARE.items():
        deck += [(f"Comparison, {title} — {TITLE_DS[ds]} — test — Teacher", fig_dir / "compare" / group / f"{ds}.png") for ds in DATASETS]
    deck += [(f"Ten-crop, sweep arms — {TITLE_DS[ds]} — test — Teacher", fig_dir / "sweep_all" / f"tencrop_{ds}.png") for ds in DATASETS]
    deck += [(f"Center crop, sweep arms — {TITLE_DS[ds]} — test — Teacher", fig_dir / "sweep_all" / f"center_{ds}.png") for ds in DATASETS]
    return deck


# ---------------------------------------------------------------- build

def build(root, container, out_dir, skip_figs):
    t0 = time.perf_counter()
    WARNINGS.clear()
    BAD.clear()
    MANIFEST.update(inputs={}, outputs=[])
    out_dir.mkdir(parents=True, exist_ok=True)
    now = datetime.now().astimezone()
    runs = run_table(root, container)
    by_id = {r["id"]: r for r in runs}
    ev = {r["id"]: load_eval(root, r) for r in runs}
    have = [r for r in runs if ev[r["id"]]]
    tabs = {r["id"]: cycle_table(ev[r["id"]]) for r in have}
    picks = {i: pick(t) for i, t in tabs.items()}
    for i, (done, _, _) in picks.items():
        if len(done) < len(tabs[i]):
            warn(f"{i}: cycle(s) {sorted(set(tabs[i]) - set(done))} not completely evaluated yet, left out of the summaries")
    probes = load_probes(root)
    ten, ssl = tencrop(root, runs), ssl_epochs(root)
    pairs = tencrop_pairs(ten, tabs, by_id, offline_files(root, runs))

    sheets = []

    def add(name, head, rows, fmt=None):
        sheets.append((name, head, rows, fmt or {}))

    add("README", ["README"], [[line] for line in readme_lines(root, [r["id"] for r in have], now)])
    rrows = []
    for r in runs:
        rc, (done, last, _) = r["recipe"], picks.get(r["id"], ([], None, None))
        rrows.append([r["id"], r["label"], r["group"], r["dir"], len(done), done[0] if done else None, last,
                      *[numish(rc[k]) if rc else None for k in RECIPE_KEYS], rc and rc["init"], rc and rc["init_tag"], rc and rc["weights"],
                      status_of(rc, ev[r["id"]], last), cycle0_kind(root, r, ev[r["id"]], ten), input_size_of(root, r), *[numish(rc[k]) if rc and rc.get(k) else None for k in SEED_COLS]])
    add("runs", ["id", "label", "group", "run_dir", "cycles_evaluated", "first_cycle", "last_cycle", *RECIPE_KEYS, "init", "init_tag", "weights", "status", "cycle0", "input_size", *SEED_COLS], rrows)
    ledger = [i for i in (*EARLY, *with_extras(SWEEP, by_id)) if i in tabs]  # the PDF ledger's runs in its order (tables_html); probe_a, probe_b and a non-adamw extra arm are in the summary sheets only
    status = {}
    for i in ledger:
        last, st = picks[i][1], status_of(by_id[i]["recipe"], ev[i], picks[i][1])  # the runs sheet's status; the ledger adds the cycle of an unfinished run
        status[i] = f"in progress (c{last} of {numish(by_id[i]['recipe']['cycles'])})" if st == "in progress" else st
    sheets.extend(ledger_sheets(tabs, picks, ledger, status))
    head = ["run_id", "label", "split", "cycle", "epoch", "evaluation_point", "dataset", "focused", "end_of_cycle", "student", "teacher"]
    is_m = lambda m: m == "mAUC"
    is_c = lambda m: m.startswith("AUC_")
    add("cyclic_mAUC", head, [x for r in have for sp in ("test", "validation") for x in eval_rows(r, ev[r["id"]], sp, is_m, False)], dict.fromkeys(head[-2:], AUC_FMT))
    chead = [*head[:7], "class", *head[7:]]
    add("cyclic_classAUC", chead, [x for r in have for x in eval_rows(r, ev[r["id"]], "test", is_c, True)], dict.fromkeys(chead[-2:], AUC_FMT))
    vrows = []
    for r in have:
        rows = [[r["id"], x["cycle"], x["epoch"], x["point"], ds, x["student"], x["teacher"]]
                for ds in DATASETS for x in ev[r["id"]].get(("validation", ds), (None, []))[1] if x["metric"] == "validation_loss"]
        vrows += sorted(rows, key=lambda x: order(x[1], x[3], x[4]))
    add("validation_loss", ["run_id", "cycle", "epoch", "evaluation_point", "dataset", "student", "teacher"], vrows, {"student": LOSS_FMT, "teacher": LOSS_FMT})
    for r in have:
        h, rows = run_sheet(tabs[r["id"]])
        add(r["id"], h, rows, dict.fromkeys(h[2:], AUC_FMT))

    def summary(which):
        rows = []
        for r in have:
            c = picks[r["id"]][which]
            if c is not None:
                t = tabs[r["id"]][c]
                rows.append([r["id"], r["label"], c, *end_values(t), pct(mean_end(t, 0))])
        return rows

    for name, which, col in (("summary_last", 1, "last_cycle"), ("summary_best_mean", 2, "best_cycle")):
        h = ["id", "label", col, *END_COLS, "T_mean_end"]
        add(name, h, summary(which), dict.fromkeys(h[3:], AUC_FMT))
    bhead, brows = ["id", "label", *[f"{SHORT[d]}_{x}" for d in DATASETS for x in ("best", "cycle")]], []
    for r in have:
        cells = []
        for d in DATASETS:
            v, c = best_of(tabs[r["id"]], d)
            cells += [pct(v), c]
        brows.append([r["id"], r["label"], *cells])
    add("summary_best_per_dataset", bhead, brows, {h: AUC_FMT for h in bhead if h.endswith("_best")})
    phead, prows, per_run = probe_sheets(probes)
    add("probes", phead, prows, {h: AUC_FMT for h in phead if h.endswith(("_mean", "_std"))})
    add("probes_per_run", ["tag", "dataset", "run", "seed", "best_epoch", "test_mAUC"], per_run, {"test_mAUC": AUC_FMT})
    add("offline_eval", ["run_id", "cycle", "dataset", "model", "metric", "class", "value", "protocol"], ten, {"value": AUC_FMT})
    thead = ["run_id", "label", "cycle", "dataset", "center_teacher", "tencrop_teacher", "delta_teacher", "center_student", "tencrop_student", "delta_student", "center_source"]
    add("tencrop_vs_center", thead, pairs, dict.fromkeys(thead[4:10], AUC_FMT))
    sfmt = dict.fromkeys((*LOSS_COLS, "teacher_temp_end", "step_time_mean"), LOSS_FMT)
    add("ssl_losses_epoch", SSL_HEADER, [[d[h] for h in SSL_HEADER] for d in ssl], {**sfmt, "lr_end": "0.00E+00"})
    for p, n in BAD.items():
        if n:
            warn(f"{p}: {n} unreadable value(s) left empty")

    if bad := ledger_disagreements(sheets, ledger, tabs):
        raise AssertionError(f"the ledger sheets disagree with the summary sheets: {bad[:5]}")

    xlsx, frag = out_dir / XLSX_NAME, out_dir / "tables.html"
    write_workbook(xlsx, sheets)
    wrote(xlsx)
    frag.write_text(tables_html(tabs, picks, probes, pairs, with_extras(SWEEP, by_id)), encoding="utf-8")
    wrote(frag)
    slides = 0
    if not skip_figs:
        fig_dir = out_dir / "figs"
        shutil.rmtree(fig_dir, ignore_errors=True)  # only this build's figures may reach the deck and the count
        full = [r for r in have if len(ev[r["id"]]) == N_FILES]  # plot_cyclic_run.py needs all eight CSVs
        for r in have:
            if r not in full:
                warn(f"{r['id']}: figures skipped, evaluation files incomplete")
        with tempfile.TemporaryDirectory(dir=SCRATCH if SCRATCH.is_dir() else None) as tmp:
            cyclic_figs(root, full, fig_dir, Path(tmp))
            compare_figs(root, by_id, {r["id"] for r in full}, fig_dir, Path(tmp))
        pages = probe_pages(root, probes, fig_dir, now.strftime("%Y-%m-%d")) if probes else []
        fig_trajectory(probes, fig_dir / "ssl" / "probe_trajectory.png")
        for name, title in LOSS_FIGS:
            fig_loss(ssl, name, title, fig_dir / "ssl" / f"losses_{name}.png")
        sweep_all_figs(ten, tabs, picks, by_id, fig_dir / "sweep_all")
        deck = deck_entries(fig_dir, tabs, by_id, pages)
        slides = make_deck([(t, p) for t, p in deck if any(q.is_file() for q in ([p] if isinstance(p, Path) else p))], out_dir / PPTX_NAME)
    man = out_dir / "build_manifest.json"
    wrote(man)
    man.write_text(json.dumps({"built": now.isoformat(timespec="seconds"), "root": str(root), **MANIFEST}, indent=1), encoding="utf-8")
    return {"out_dir": out_dir, "sheets": [s[0] for s in sheets], "slides": slides, "seconds": time.perf_counter() - t0}


# ---------------------------------------------------------------- selftest and CLI

def synth(ri, cycle, pi, di, k):
    """Unique value per (run, cycle, evaluation point, dataset, model: 0 student / 1 teacher). Run 0 is best at its first cycle; run 1 ties at both (best = earliest)."""
    return 0.5 + 0.05 * ri + (0.1 if ri == 0 and cycle == 2 else 0.0) + 0.01 * pi + 0.001 * di + 0.02 * k


def synth_mirror(root, container):
    import plot_linear_probes as P

    cycles = {"all4_cyclic_dinov3_sweep_S2_seed_100": (2, 4), "all4_cyclic_dinov3_frozen_long_seed_100": (5, 10)}
    for ri, (rdir, cyc) in enumerate(cycles.items()):
        for sp, fn in SPLIT_FILES:
            for di, ds in enumerate(DATASETS):
                p = root / rdir / "evaluation" / ds / fn
                p.parent.mkdir(parents=True, exist_ok=True)
                with open(p, "w", newline="") as f:  # csv.writer writes CRLF, like the real files
                    w = csv.writer(f)
                    w.writerow(["cycle", "epoch", "evaluation_point", "metric", "student", "teacher"])
                    for c in cyc:
                        for pi, pt in enumerate(POINTS):
                            for m in ("AUC_x", "mAUC", *(("validation_loss",) if sp == "validation" and ds != "VinDrCXR" else ())):
                                w.writerow([c, c - 1, pt, m, synth(ri, c, pi, di, 0), synth(ri, c, pi, di, 1)])
    (container / "sweep").mkdir(parents=True)
    head = "id opt lr wd layer_decay patch_embed_mult warmup cycles model init weights normalization".split()
    row = "S2 adamw 7e-5 0.01 0.75 1.0 2 4 vit_base_dinov3 dinov3 /w/dinov3_vitb16_pretrain_x.pth imagenet".split()
    (container / "sweep" / "wave1.tsv").write_text("\t".join(head) + "\n" + "\t".join(row) + "\n")
    P.make_synthetic(str(root / "results"), tags=("stock_dinov3",))
    with open(root / "cycle_0002_all4_cyclic_dinov3_sweep_S2_seed_100.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["dataset", "model", "metric", "class", "value"])
        for ds in DATASETS:
            for model in ("student", "teacher"):
                w.writerow([ds, model, "mAUC", "", 0.9])
    (root / "ssl_dinov3_cxr_main_seed_100").mkdir()
    with open(root / "ssl_dinov3_cxr_main_seed_100" / "ssl_losses.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["step", "epoch", *LOSS_COLS, "lr", "wd", "momentum", "teacher_temp", "gram_weight", "step_time", "data_time"])
        for step in range(7):
            w.writerow([step, step // 3, *[float(step)] * len(LOSS_COLS), 0.001 * step, 0.04, 0.999, 0.04, 2.0, 1.0, 0.1])


def off_value(protocol, cycle, di, k):
    """The synthetic offline mAUC (a fraction) per protocol, cycle, dataset and model (0 student / 1 teacher)."""
    return (0.80 if protocol == "tencrop" else 0.60) + 0.01 * cycle + 0.001 * di + 0.002 * k


def init_value(di, k, init=0.99):
    """The synthetic cycle-0 mAUC (a fraction) per dataset and model (0 student / 1 teacher); the two differ, as two separate model builds do."""
    return init - (0.05 if k == 0 else 0.0) + 0.001 * di


def write_eval(root, rdir, ri, cycles, init=None, points=POINTS):
    """The evaluation CSVs of one run (four datasets, test and validation); init = the cycle-0 teacher value (a fraction; init_value gives both models), else no cycle-0 rows;
    points = the evaluation points each cycle has (a run partway through a cycle has the first ones only)."""
    for sp, fn in SPLIT_FILES:
        for di, ds in enumerate(DATASETS):
            p = root / rdir / "evaluation" / ds / fn
            p.parent.mkdir(parents=True, exist_ok=True)
            metrics = ("AUC_x", "mAUC", *(("validation_loss",) if sp == "validation" and ds != "VinDrCXR" else ()))
            with open(p, "w", newline="") as f:
                w = csv.writer(f)
                w.writerow(["cycle", "epoch", "evaluation_point", "metric", "student", "teacher"])
                for m in metrics if init is not None else ():
                    w.writerow([0, 0, INIT, m, init_value(di, 0, init), init_value(di, 1, init)])
                for c in cycles:
                    for pi, pt in enumerate(points):
                        for m in metrics:
                            w.writerow([c, c - 1, pt, m, synth(ri, c, pi, di, 0), synth(ri, c, pi, di, 1)])


def synth_fresh(root, container):
    """A mirror of two runs that have only just started: S3 has its cycle-0 init rows and nothing else, S4 is partway through cycle 1 (init, after_VinDrCXR, after_ChestXray14).
    Returns their run directories."""
    s3, s4 = "all4_cyclic_dinov3_sweep_S3_seed_100", "all4_cyclic_dinov3_sweep_S4_seed_100"
    write_eval(root, s3, 2, (), 0.97)
    write_eval(root, s4, 3, (1,), 0.96, POINTS[:2])
    (container / "sweep").mkdir(parents=True)
    head = "id opt lr wd layer_decay patch_embed_mult warmup cycles model init weights normalization".split()
    rows = [f"{i} adamw 7e-5 0.01 0.75 1.0 2 4 vit_base_dinov3 dinov3 /w/dinov3_vitb16_pretrain_x.pth imagenet".split() for i in ("S3", "S4")]
    (container / "sweep" / "wave1.tsv").write_text("\n".join("\t".join(r) for r in (head, *rows)) + "\n")
    return s3, s4


def synth_cycle0(root, container, init=0.99):
    """A two-run mirror. S2 is evaluated at cycle 0 (the init point; here its best cycle) and at cycles 1 (epoch 0 too), 2 and 4, and has the offline files: ten-crop for
    cycles 0-4, center crop for cycles 0-3 (cycle 4 has none). D-long has no cycle-0 rows."""
    s2, dl = "all4_cyclic_dinov3_sweep_S2_seed_100", "all4_cyclic_dinov3_frozen_long_seed_100"
    write_eval(root, s2, 0, (1, 2, 4), init)
    write_eval(root, dl, 1, (5, 10))
    (container / "sweep").mkdir(parents=True)
    head = "id opt lr wd layer_decay patch_embed_mult warmup cycles model init weights normalization".split()
    row = "S2 adamw 7e-5 0.01 0.75 1.0 2 4 vit_base_dinov3 dinov3 /w/dinov3_vitb16_pretrain_x.pth imagenet".split()
    (container / "sweep" / "wave1.tsv").write_text("\t".join(head) + "\n" + "\t".join(row) + "\n")
    for suffix, protocol, cycles in (("", "tencrop", range(5)), (".centercrop", "centercrop", range(4))):
        for c in cycles:
            with open(root / f"cycle_{c:04d}_{s2}{suffix}.csv", "w", newline="") as f:
                w = csv.writer(f)
                w.writerow(["dataset", "model", "metric", "class", "value"])
                for di, ds in enumerate(DATASETS):
                    for k, model in enumerate(("student", "teacher")):
                        w.writerow([ds, model, "mAUC", "", off_value(protocol, c, di, k)])
    return s2, dl


def selftest_cycle0(tmp):
    """Cycle 0, the offline protocols, the hollow markers and the plot scripts on a synthetic mirror (cycle 0 is S2's best cycle there: it must be picked)."""
    from types import SimpleNamespace

    from openpyxl import load_workbook
    import compare_cyclic_runs as K
    import plot_cyclic_run as C

    root, container, out = Path(tmp) / "c0", Path(tmp) / "c0c", Path(tmp) / "c0out"
    s2, dl = synth_cycle0(root, container)
    build(root, container, out, skip_figs=True)
    xlsx = out / XLSX_NAME
    wb = load_workbook(xlsx)

    def rows(name):
        ws = wb[name]
        head = [c.value for c in ws[1]]
        return [dict(zip(head, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]

    # 1. cycle 0 in the sheets and in the pick
    init = lambda di: pct(init_value(di, 1))  # the teacher's init value
    per_run = rows("S2")
    assert [(r["cycle"], r["epoch"]) for r in per_run] == [(0, 0), (1, 0), (2, 1), (4, 3)], per_run  # cycle 1 is epoch 0 too: the rows are told apart by cycle, never by epoch
    assert all(per_run[0][f"{w}_{SHORT[d]}_{kind}"] == pct(init_value(di, int(w == "T"))) for di, d in enumerate(DATASETS) for w in "TS" for kind in ("end", "focused")), per_run[0]  # end-of-cycle AND focused, a value of its own per model
    assert all(per_run[0][f"T_{SHORT[d]}_end"] != per_run[0][f"S_{SHORT[d]}_end"] for d in DATASETS), per_run[0]  # the student and the teacher differ at init: no equality is assumed
    assert per_run[1]["T_CheXpert_end"] == pct(synth(0, 1, 3, 3, 1)) and per_run[1]["T_VinDr_focused"] == pct(synth(0, 1, 0, 0, 1)) and per_run[2]["T_CheXpert_end"] == pct(synth(0, 2, 3, 3, 1)) and per_run[2]["T_VinDr_focused"] == pct(synth(0, 2, 0, 0, 1)), per_run  # the in-cycle points are untouched
    m0 = [r for r in rows("cyclic_mAUC") if r["evaluation_point"] == INIT]
    assert len(m0) == 2 * 4 and all((r["cycle"], r["epoch"], r["focused"], r["end_of_cycle"]) == (0, 0, True, True) for r in m0), m0  # test and validation
    assert all((r["student"], r["teacher"]) == tuple(pct(init_value(DATASETS.index(r["dataset"]), k)) for k in (0, 1)) and r["student"] != r["teacher"] for r in m0), m0
    c0 = [r for r in rows("cyclic_classAUC") if r["evaluation_point"] == INIT]
    assert len(c0) == 4 and all(r["focused"] is True and r["end_of_cycle"] is True for r in c0), c0
    assert {r["id"]: r["best_cycle"] for r in rows("summary_best_mean")} == {"S2": 0, "D-long": 5}, rows("summary_best_mean")  # cycle 0 wins on S2: picked, not an error
    assert {r["id"]: r["last_cycle"] for r in rows("summary_last")} == {"S2": 4, "D-long": 10}
    assert [r[f"{SHORT[d]}_cycle"] for r in rows("summary_best_per_dataset") if r["id"] == "S2" for d in DATASETS] == [0] * 4
    runs_ = {r["id"]: r for r in rows("runs")}
    assert (runs_["S2"]["first_cycle"], runs_["S2"]["cycles_evaluated"], runs_["S2"]["cycle0"]) == (0, 4, "init") and runs_["D-long"]["cycle0"] is None, runs_
    assert not any("not completely evaluated" in w for w in WARNINGS), WARNINGS
    inputs = lambda: json.loads((out / "build_manifest.json").read_text(encoding="utf-8"))["inputs"]
    assert not any(p.endswith("reconstructed_from_seed.json") for p in inputs()), inputs()  # no marker yet: none listed
    d0 = root / s2 / "models" / "weights" / "epoch_0000"
    d0.mkdir(parents=True)
    (d0 / "reconstructed_from_seed.json").write_text("{}")
    build(root, container, out, skip_figs=True)
    wb = load_workbook(xlsx)
    assert {r["id"]: r["cycle0"] for r in rows("runs")} == {"S2": "reconstructed init", "D-long": None}, rows("runs")  # the json marks the cycle-0 point as reconstructed
    assert inputs().get(str(d0 / "reconstructed_from_seed.json"), {}).get("size") == 2, inputs()  # ...and it is an input of the build: build_manifest.json lists it, with its size and mtime

    # 2. the offline files: one sheet, a protocol column; pairing with the offline center crop where it exists
    assert "tencrop" not in wb.sheetnames and wb.sheetnames[wb.sheetnames.index("offline_eval") - 1] == "probes_per_run", wb.sheetnames
    off = rows("offline_eval")
    assert [c.value for c in wb["offline_eval"][1]][-1] == "protocol" and Counter(r["protocol"] for r in off) == {"tencrop": 5 * 8, "centercrop": 4 * 8}, Counter(r["protocol"] for r in off)
    tv = rows("tencrop_vs_center")
    assert len(tv) == 5 * 4 and [r["cycle"] for r in tv if r["dataset"] == "VinDrCXR"] == [0, 1, 2, 3, 4], len(tv)
    for r in tv:
        c, di = r["cycle"], DATASETS.index(r["dataset"])
        assert r["center_source"] == ("offline" if c <= 3 else "in-training"), r  # the offline center crop where there is one, else the in-training one (cycle 4)
        for k, m in enumerate(("student", "teacher")):
            ctr = pct(off_value("centercrop", c, di, k)) if c <= 3 else pct(synth(0, 4, 3, di, k))
            ten = pct(off_value("tencrop", c, di, k))
            assert (r[f"center_{m}"], r[f"tencrop_{m}"], r[f"delta_{m}"]) == (ctr, ten, round(ten - ctr, 2)), r
    assert not any("ten-crop cycle" in w for w in WARNINGS), WARNINGS
    frag = (out / "tables.html").read_text(encoding="utf-8")

    def tc_row(c, pairs):  # S2's row of the ledger table: (center, ten-crop) teacher for the four datasets, one decimal
        cells = "".join(f'<td class="{"r sep" if di == j == 0 else "r"}">{f1(v)}</td>' for di, pair in enumerate(pairs) for j, v in enumerate(pair))
        return f'<tr><td class="id">S2</td><td class="cyc">c{c}</td>{cells}</tr>'

    assert tc_row(1, [(off_value("centercrop", 1, di, 1), off_value("tencrop", 1, di, 1)) for di in range(4)]) in frag  # the same rule as the sheet: offline center crop
    assert tc_row(4, [(synth(0, 4, 3, di, 1), off_value("tencrop", 4, di, 1)) for di in range(4)]) in frag  # cycle 4 has none: the in-training center crop
    readme = [r["README"] for r in rows("README")]
    for start in ("Cycle 0 is the initial model before any training:", "Two offline protocols evaluate the saved weights", "Sweep-wide center-crop figures draw the in-training",
                  "offline_eval (renamed from tencrop):"):
        assert sum(l.startswith(start) for l in readme) == 1, (start, readme)
    assert not any(l.startswith("tencrop:") for l in readme) and any("hollow = offline evaluation of the saved weights" in l for l in readme), readme
    dl0 = root / dl / "models" / "weights" / "epoch_0000"  # a reconstructed init: no cycle-0 rows in the evaluation CSVs, the json says the weights were rebuilt from the seed, the point is in cycle_0000 offline files
    dl0.mkdir(parents=True)
    (dl0 / "reconstructed_from_seed.json").write_text("{}")
    for suffix, protocol in (("", "tencrop"), (".centercrop", "centercrop")):
        with open(root / f"cycle_0000_{dl}{suffix}.csv", "w", newline="") as f:
            w = csv.writer(f)
            w.writerow(["dataset", "model", "metric", "class", "value"])
            for di, ds in enumerate(DATASETS):
                for k, model in enumerate(("student", "teacher")):
                    w.writerow([ds, model, "mAUC", "", off_value(protocol, 0, di, k)])
    build(root, container, out, skip_figs=True)
    wb = load_workbook(xlsx)
    runs_ = {r["id"]: r for r in rows("runs")}
    assert (runs_["D-long"]["cycle0"], runs_["D-long"]["first_cycle"]) == ("reconstructed init", 5), runs_["D-long"]  # its cycle-0 point exists only offline
    assert {str(d0 / "reconstructed_from_seed.json"), str(dl0 / "reconstructed_from_seed.json")} <= set(inputs()), inputs()  # both markers are listed
    assert {(r["run_id"], r["protocol"]) for r in rows("offline_eval") if r["run_id"] == "D-long"} == {("D-long", "tencrop"), ("D-long", "centercrop")}
    assert [(r["cycle"], r["center_source"]) for r in rows("tencrop_vs_center") if r["run_id"] == "D-long"] == [(0, "offline")] * 4, rows("tencrop_vs_center")

    # 3. the plot scripts: the init point at x = 0, the axis from the first point (0), whole-cycle ticks; a run without cycle-0 rows is as before
    def spying(store, n):
        def spy(fig, path, dpi):
            ax = fig.axes[0]
            store[path.parts[-n:]] = ([(list(ln.get_xdata()), list(ln.get_ydata())) for ln in ax.get_lines()], tuple(ax.get_xlim()), ax.get_xticks().tolist())
            C.plt.close(fig)

        return spy

    real_c, real_k = C._save_figure, K._save_figure
    cargs = SimpleNamespace(run_directory=root / s2, output_root=Path(tmp) / "figsC", metric="mAUC", format="png", dpi=30, y_min=None, y_max=None)
    seen_c, seen_d, seen_k = {}, {}, {}
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            C._save_figure = spying(seen_c, 3)
            C.create_plots(cargs)
            C._save_figure = spying(seen_d, 3)
            C.create_plots(SimpleNamespace(**{**vars(cargs), "run_directory": root / dl}))
            K._save_figure = spying(seen_k, 2)
            K.create_plots(SimpleNamespace(run_directories=[root / s2, root / dl], metric="mAUC", output_root=Path(tmp) / "figsK", split="test", models="teacher", colors=[], labels=[],
                                           dpi=30, format="png", y_min=None, y_max=None, full_length=True))
    finally:
        C._save_figure, K._save_figure = real_c, real_k
    lines, xlim, ticks = seen_c[("focused_unfocused", "test", "VinDrCXR.png")]
    assert len(lines) == 4 and all(xs[0] == 0.0 for xs, _ in lines) and xlim[0] == 0.0 and all(t == int(t) for t in ticks), (xlim, ticks)  # every curve starts on the y axis
    assert lines[0][0][:6] == [0.0, 0.25, 0.5, 0.75, 1.0, 1.25], lines[0]  # init at 0, then cycle 1 (epoch 0 too, not mistaken for init) at its quarter positions, then cycle 2
    assert abs(lines[0][1][0] - 100 * init_value(0, 0)) < 1e-9 and abs(lines[1][1][0] - 100 * init_value(0, 1)) < 1e-9 and lines[0][1][0] != lines[1][1][0], (lines[0][1][0], lines[1][1][0])  # both models drawn from their own init value
    assert lines[2][0] == lines[3][0] == [0.0, 0.25, 1.25, 3.25] and (lines[2][1][0], lines[3][1][0]) == (lines[0][1][0], lines[1][1][0]), lines[2:]  # the focused curve: the init point, then the dataset's own point in each evaluated cycle
    dlines, dxlim, _ = seen_d[("focused_unfocused", "test", "VinDrCXR.png")]
    assert all(xs[0] == 4.25 for xs, _ in dlines) and dxlim[0] == 4.25, (dlines[0][0][:2], dxlim)  # no cycle-0 rows: from the first evaluated point, as before
    klines, kxlim, kticks = seen_k[("test", "VinDrCXR.png")]
    assert [xs[0] for xs, _ in klines] == [0.0, 4.25] and kxlim[0] == 0.0 and all(t == int(t) for t in kticks), (klines, kxlim)  # compare_cyclic_runs: S2 from the init point, D-long from its first cycle

    # 4. the sweep-wide figures: in-training values solid, offline center-crop values hollow where there is no in-training value; the ten-crop curve is the ten-crop protocol only
    runs = run_table(root, container)
    tabs = {r["id"]: cycle_table(files) for r in runs if (files := load_eval(root, r))}
    picks, ten, by_id = {i: pick(t) for i, t in tabs.items()}, tencrop(root, runs), {r["id"]: r for r in runs}
    di = DATASETS.index("MIMIC")
    tcs, ccs, hollow = sweep_series(ten, tabs, picks, "MIMIC")
    assert tcs == {"S2": [(c, pct(off_value("tencrop", c, di, 1))) for c in range(5)]}, tcs
    assert ccs == {"S2": [(0, init(di)), (1, pct(synth(0, 1, 3, di, 1))), (2, pct(synth(0, 2, 3, di, 1))), (3, pct(off_value("centercrop", 3, di, 1))), (4, pct(synth(0, 4, 3, di, 1)))]}, ccs
    assert hollow == {"S2": {3}}, hollow  # cycles 0 (the init point), 1, 2 and 4 are in-training values, only cycle 3 is offline
    tabs2 = {"S2": {c: v for c, v in tabs["S2"].items() if c != 0}}
    _, cc2, hol2 = sweep_series(ten, tabs2, {"S2": pick(tabs2["S2"])}, "MIMIC")
    assert [c for c, _ in cc2["S2"]] == [0, 1, 2, 3, 4] and hol2 == {"S2": {0, 3}} and cc2["S2"][0][1] == pct(off_value("centercrop", 0, di, 1)), (cc2, hol2)  # no init in training: cycle 0 is offline
    part = {c: ({**v, "end": {d: tv for d, tv in v["end"].items() if d != DATASETS[0]}} if c in (2, 4) else v) for c, v in tabs["S2"].items()}  # cycles 2 and 4 lack VinDrCXR (the mirror copies the evaluation files one at a time): incomplete, MIMIC's own value is there
    assert pick(part)[0] == [0, 1], pick(part)
    _, cc3, hol3 = sweep_series(ten, {"S2": part}, {"S2": pick(part)}, "MIMIC")
    assert (cc3, hol3) == (ccs, hollow), (cc3, hol3)  # the dataset's own value decides, not the complete cycles: cycle 4 (no offline center crop) is not dropped, cycle 2 (an offline one) is not replaced by a hollow point
    seen, real_save = {}, save_fig

    def spy_fig(plt, fig, path):
        ax = fig.axes[0]
        seen[path.name] = ([(list(ln.get_xdata()), ln.get_linestyle(), ln.get_markerfacecolor(), ln.get_markeredgecolor(), ln.get_color()) for ln in ax.get_lines()],
                           [t.get_text() for t in ax.get_legend().get_texts()], tuple(ax.get_xlim()))
        real_save(plt, fig, path)

    globals()["save_fig"] = spy_fig
    try:
        sweep_all_figs(ten, tabs, picks, by_id, Path(tmp) / "sw")
        mlines, mlegend, mxlim = seen["center_MIMIC.png"]
        assert len(mlines) == 2 and mlines[0][0] == [0, 1, 2, 3, 4] and mlines[0][4] == ARM_COLORS["S2"], mlines  # one line through every point...
        assert mlines[1][:2] == ([3], "None") and mlines[1][2:4] == ("white", ARM_COLORS["S2"]), mlines[1]  # ...and hollow markers of the arm's colour at the offline cycles
        assert mlegend.count(CC_NOTE) == 1 and mlegend[-1] == CC_NOTE and mxlim[0] == 0, (mlegend, mxlim)  # the one legend note; the axis starts at 0
        tlines, tlegend, txlim = seen["tencrop_MIMIC.png"]
        assert len(tlines) == 1 and tlines[0][0] == [0, 1, 2, 3, 4] and CC_NOTE not in tlegend and txlim[0] == 0, (tlines, tlegend)  # ten-crop: the same form, now dense
        seen.clear()
        sweep_all_figs([r for r in ten if r[7] == "tencrop"], tabs, picks, by_id, Path(tmp) / "sw2")
        nlines, nlegend, _ = seen["center_MIMIC.png"]
        assert len(nlines) == 1 and CC_NOTE not in nlegend, (nlines, nlegend)  # no offline point: no hollow marker and no note
        seen.clear()
        sweep_all_figs(ten, {}, {}, by_id, Path(tmp) / "sw3")  # nothing evaluated in training: every point of the center curve is offline
        alines, alegend, _ = seen["center_MIMIC.png"]
        assert alines[0][0] == alines[1][0] == [0, 1, 2, 3] and CC_NOTE in alegend, alines
    finally:
        globals()["save_fig"] = real_save

    # 5. runs that have only just started: S3 has its cycle-0 init rows and nothing else, S4 is partway through cycle 1 (init, after_VinDrCXR, after_ChestXray14)
    froot, fcont, fout = Path(tmp) / "fresh", Path(tmp) / "freshc", Path(tmp) / "freshout"
    d3, d4 = synth_fresh(froot, fcont)
    build(froot, fcont, fout, skip_figs=True)  # the workbook takes both as they are
    wb = load_workbook(fout / XLSX_NAME)
    fr = {r["id"]: r for r in rows("runs")}
    assert (fr["S3"]["cycle0"], fr["S3"]["first_cycle"], fr["S3"]["last_cycle"], fr["S4"]["cycle0"]) == ("init", 0, 0, "init"), fr  # cycle 0 alone is a complete cycle
    assert [(r["cycle"], r["T_VinDr_end"]) for r in rows("S3")] == [(0, pct(init_value(0, 1, 0.97)))], rows("S3")
    fruns = run_table(froot, fcont)
    ffig = fout / "figs"
    cyclic_figs(froot, fruns, ffig, Path(tmp) / "fcyc")  # the per-run figure script, once per run as the build runs it: it exits non-zero on a run it cannot draw, and the build stops there
    compare_figs(froot, {r["id"]: r for r in fruns}, {"S3", "S4"}, ffig, Path(tmp) / "fcmp")
    assert all((ffig / "cyclic" / rid / f"{kind}_{split}_{ds}.png").is_file() for rid in ("S3", "S4") for kind, split in FIG_KINDS for ds in DATASETS)
    assert all((ffig / "compare" / "adamw" / f"{ds}.png").is_file() for ds in DATASETS), sorted(p.name for p in ffig.rglob("*.png"))

    def spying_m(store, n):  # like spying, plus the marker of each line
        def spy(fig, path, dpi):
            store[path.parts[-n:]] = [(list(ln.get_xdata()), list(ln.get_ydata()), ln.get_marker()) for ln in fig.axes[0].get_lines()]
            C.plt.close(fig)

        return spy

    seen_3, seen_4, seen_f, said = {}, {}, {}, io.StringIO()
    try:
        with contextlib.redirect_stdout(said):
            C._save_figure = spying_m(seen_3, 3)
            C.create_plots(SimpleNamespace(**{**vars(cargs), "run_directory": froot / d3, "output_root": Path(tmp) / "f3"}))
            C._save_figure = spying_m(seen_4, 3)
            C.create_plots(SimpleNamespace(**{**vars(cargs), "run_directory": froot / d4, "output_root": Path(tmp) / "f4"}))
            K._save_figure = spying_m(seen_f, 2)
            K.create_plots(SimpleNamespace(run_directories=[froot / d4, froot / d3], metric="mAUC", output_root=Path(tmp) / "ff", split="test", models="teacher", colors=[], labels=[],
                                           dpi=30, format="png", y_min=None, y_max=None, full_length=True))
    finally:
        C._save_figure, K._save_figure = real_c, real_k
    assert "Evaluation order: VinDrCXR -> ChestXray14\nNot evaluated yet: CheXpert, MIMIC\n" in said.getvalue() and "Evaluation order: (none yet: the init point only)\n" in said.getvalue(), said.getvalue()  # the console lists the tasks seen, and says which are not evaluated yet
    assert len(seen_3) == 24 and all(ls and all(xs == [0.0] and len(ys) == 1 and m == "o" for xs, ys, m in ls) for ls in seen_3.values()), seen_3  # S3: every curve of every figure is the init point alone, at x = 0, drawn as a marker
    assert {round(ys[0], 6) for _, ys, _ in seen_3[("focused_unfocused", "test", "VinDrCXR.png")]} == {round(100 * init_value(0, k, 0.97), 6) for k in (0, 1)}, seen_3  # each model at its own init value
    assert seen_4[("focused", "test", "MIMIC.png")] and all(xs == [0.0] and m == "o" for xs, _, m in seen_4[("focused", "test", "MIMIC.png")]), seen_4  # MIMIC has no after_MIMIC row yet: its focused curve is the init point, not nothing
    assert all(xs == [0.0, 0.5] for xs, _, _ in seen_4[("focused", "test", "ChestXray14.png")]), seen_4  # init, then its own point: the second of the cycle's four tasks, at 0.5
    assert [xs for xs, _, _ in seen_f[("test", "VinDrCXR.png")]] == [[0.0, 0.25, 0.5], [0.0]] and seen_f[("test", "VinDrCXR.png")][1][2] == "o", seen_f  # compare_cyclic_runs: S4 from its init at the full cycle's positions, S3 as its init marker
    fd = C._discover_run(froot / d4, "mAUC")
    focus = lambda ds: [r.evaluation_point for r in C._focused_records(fd.records[(ds, "test")], ds, fd.task_order)]
    assert fd.task_order[:2] == ("VinDrCXR", "ChestXray14") and set(fd.task_order) == set(DATASETS) and len(fd.task_order) == 4, fd.task_order  # the two tasks seen, then the other two datasets: the cycle's full count
    assert focus("ChestXray14") == ["init", "after_ChestXray14"] and focus("MIMIC") == focus("CheXpert") == ["init"], (fd.task_order, focus("MIMIC"))
    odd = Path(tmp) / "odd"  # init rows in some files, an after_ row in another: not a fresh run, so still the error
    shutil.copytree(froot / d3, odd)
    for fn in ("test_performance.csv", "val_performance.csv"):
        (odd / "evaluation" / "ChestXray14" / fn).write_text("cycle,epoch,evaluation_point,metric,student,teacher\n1,0,after_VinDrCXR,mAUC,0.5,0.6\n")
    try:
        C._discover_run(odd, "mAUC")
    except ValueError as e:
        assert "Could not infer evaluation order" in str(e), e
    else:
        raise AssertionError("a run with an after_ row in one file and only init rows in the others was accepted")
    # the cycle's task count is the full one (a task per dataset directory) however few evaluations have arrived: the first tasks stay at 0.25 and 0.5
    for k in (1, 2, 3, 4):
        write_eval(froot, f"part{k}", 4, (1,), 0.96, POINTS[:k])
    write_eval(froot, "joint", 5, (1,), 0.95, ("after_joint",))  # a concurrent run: one shared event per cycle, no dataset
    write_eval(froot, "evalonly", 6, (1, 2), None, POINTS[:2])  # two tasks in each of two cycles, though four datasets are evaluated

    def laid_out(name):  # the run, and the x positions of its VinDrCXR test teacher curve
        d = C._discover_run(froot / name, "mAUC")
        return d, C._series(d.records[("VinDrCXR", "test")], "teacher", {t: n for n, t in enumerate(d.task_order)}, len(d.task_order))[0]

    assert [laid_out(f"part{k}")[1] for k in (1, 2, 3, 4)] == [[0.0, 0.25], [0.0, 0.25, 0.5], [0.0, 0.25, 0.5, 0.75], [0.0, 0.25, 0.5, 0.75, 1.0]], "positions move as evaluations arrive"
    assert all(len(laid_out(f"part{k}")[0].task_order) == 4 and laid_out(f"part{k}")[0].task_order[:k] == DATASETS[:k] for k in (1, 2, 3, 4))  # the tasks seen first, in their order
    p1 = laid_out("part1")[0]
    assert [r.evaluation_point for r in C._focused_records(p1.records[("MIMIC", "test")], "MIMIC", p1.task_order)] == ["init"]  # one task seen is not a joint run: MIMIC's focused curve is its init point alone
    jd, jx = laid_out("joint")
    assert jd.task_order == ("joint",) and jx == [0.0, 1.0], (jd.task_order, jx)  # a shared event is no dataset: the order stays as seen
    ed, ex = laid_out("evalonly")
    assert ed.task_order == ("VinDrCXR", "ChestXray14") and ex == [0.5, 1.0, 1.5, 2.0], (ed.task_order, ex)  # a second cycle shows the first one whole: the order stays as seen

    # 6. the center value's source is chosen per run and cycle: an offline center-crop file that lacks a dataset leaves that dataset's cell empty, with a warning, never the in-training value
    r6, c6, o6 = Path(tmp) / "c06", Path(tmp) / "c06c", Path(tmp) / "c06out"
    s2_6, _ = synth_cycle0(r6, c6)
    cc2 = r6 / f"cycle_0002_{s2_6}.centercrop.csv"
    with open(cc2, newline="") as f:
        kept = [r for r in csv.reader(f) if r[0] != "MIMIC"]  # cycle 2's file loses MIMIC's rows
    with open(cc2, "w", newline="") as f:
        csv.writer(f).writerows(kept)
    with open(r6 / f"cycle_0003_{s2_6}.centercrop.csv", "w", newline="") as f:  # cycle 3's file has per-class rows only: it exists, and holds no mAUC row of any dataset
        w = csv.writer(f)
        w.writerow(["dataset", "model", "metric", "class", "value"])
        w.writerow(["VinDrCXR", "teacher", "AUC", "Atelectasis", 0.7])
    (r6 / f"cycle_0000_{s2_6}.centercrop.csv").write_text("")  # cycle 0's file is empty, cycle 1's holds only its header
    (r6 / f"cycle_0001_{s2_6}.centercrop.csv").write_text("dataset,model,metric,class,value\r\n")
    build(r6, c6, o6, skip_figs=True)
    wb = load_workbook(o6 / XLSX_NAME)
    tv = {(r["cycle"], r["dataset"]): r for r in rows("tencrop_vs_center")}
    for c in (0, 1):
        for ds in DATASETS:
            assert tuple(tv[(c, ds)][k] for k in ("center_source", "center_teacher", "center_student", "delta_teacher", "delta_student")) == ("offline", None, None, None, None), tv[(c, ds)]  # a file with no row is still the source: empty, never the training value
    assert [sum(f"cycle_000{c}_{s2_6}.centercrop.csv" in w for w in WARNINGS) for c in (0, 1)] == [1, 1], WARNINGS  # one warning for each of the two files, naming it
    assert not any(w.startswith(f"S2: ten-crop cycle {c} ") for c in (0, 1) for w in WARNINGS), WARNINGS  # none for the datasets
    for ds in DATASETS:
        di = DATASETS.index(ds)
        m = tv[(2, ds)]
        want = (None, None) if ds == "MIMIC" else (pct(off_value("centercrop", 2, di, 1)), pct(off_value("centercrop", 2, di, 0)))
        assert m["center_source"] == "offline" and (m["center_teacher"], m["center_student"]) == want and m["tencrop_teacher"] == pct(off_value("tencrop", 2, di, 1)), m  # MIMIC empty, the others offline
        assert (tv[(3, ds)]["center_source"], tv[(3, ds)]["center_teacher"], tv[(3, ds)]["delta_teacher"]) == ("offline", None, None), tv[(3, ds)]  # a file without mAUC rows is still the source: all empty
        assert any(all(x in w for x in ("S2", f"cycle {c}", ds, "offline center-crop")) and "end-of-cycle" not in w for c in ((2, 3) if ds == "MIMIC" else (3,)) for w in WARNINGS), (ds, WARNINGS)
    assert not any("cycle 2" in w and ds in w for ds in DATASETS if ds != "MIMIC" for w in WARNINGS), WARNINGS  # no warning for the datasets cycle 2's file has
    assert tv[(4, "MIMIC")]["center_source"] == "in-training" and tv[(4, "MIMIC")]["center_teacher"] == pct(synth(0, 4, 3, 2, 1)), tv[(4, "MIMIC")]  # no file for cycle 4: the training evaluation, as before


def selftest_input_size(tmp):
    """The sweep table's input_size column: runs.input_size from the lowest-cycle manifest.json in the mirror (224 for a manifest from before the key, an
    unreadable one left empty with a warning), else the table's value; the input-size arms' one label (run_labels.ARM_LABELS) in the workbook, the titles and
    short_run_labels; their colours."""
    import run_labels
    from openpyxl import load_workbook

    cols = "id opt lr wd layer_decay patch_embed_mult warmup cycles model init weights normalization seed eval_every eval_cycle0 input_size".split()
    row = lambda i, size: [i, "adamw", "7e-5", "0.01", "0.75", "1.0", "2", "30", "vit_base_dinov3", "dinov3", "/w/dinov3_vitb16_pretrain_x.pth", "imagenet", "100", "1", "false", size]
    root, container, out = Path(tmp) / "size", Path(tmp) / "size-container", Path(tmp) / "size-report"
    (container / "sweep").mkdir(parents=True)
    ids = (("S2", "224"), ("S18", "448"), ("S19", "512"), ("X1", "448"), ("X2", "448"))
    (container / "sweep" / "wave1.tsv").write_text("\n".join("\t".join(r) for r in (cols, *(row(i, n) for i, n in ids))) + "\n")
    for ri, (i, _) in enumerate(ids):
        write_eval(root, f"all4_cyclic_dinov3_sweep_{i}_seed_100", ri, (1, 2))
    manifests = {"S18": {1: '{"input_size": 448, "img_resize": 512}', 2: '{"input_size": 999}'},  # the lowest cycle's is read
                 "X1": {1: '{"cycle": 1}'},  # a manifest from before the key: 224, whatever the table says
                 "X2": {1: "{"}}  # unreadable: empty cell and a warning
    for i, docs in manifests.items():
        for cycle, text in docs.items():
            d = root / f"all4_cyclic_dinov3_sweep_{i}_seed_100" / "models" / "weights" / f"epoch_{cycle:04d}"
            d.mkdir(parents=True)
            (d / "manifest.json").write_text(text)
    build(root, container, out, skip_figs=True)
    ws = load_workbook(out / XLSX_NAME)["runs"]
    head = [c.value for c in ws[1]]
    got = {r[head.index("id")]: r[head.index("input_size")] for r in ws.iter_rows(min_row=2, values_only=True)}
    labels = {r[head.index("id")]: r[head.index("label")] for r in ws.iter_rows(min_row=2, values_only=True)}
    s18 = "S18 (center recipe, 448 px)"
    assert labels["S18"] == s18 == run_labels.arm_label("S18") and run_labels.short_run_labels([root / "all4_cyclic_dinov3_sweep_S18_seed_100"]) == [s18], labels
    assert labels["S19"] == "S19 (center recipe, 512 px)" and list(run_labels.ARM_LABELS) == ["S18", "S19"], labels
    assert labels["S2"] == "S2 (adamw lr 7e-5, wd 0.01, ld 0.75, pe 1.0, init stock)" and labels["X1"].startswith("X1 (adamw"), labels  # the others stay derived
    assert head[-4:] == ["input_size", *SEED_COLS] and got == {"S2": 224, "S18": 448, "S19": 512, "X1": 224, "X2": None}, (head, got)
    assert any("X2: unreadable" in w and "input_size left empty" in w for w in WARNINGS), WARNINGS
    by_id = {r["id"]: r for r in run_table(root, container)}
    center = by_id["S2"]["recipe"]
    assert [short_title(by_id[i], center) for i in ("S18", "S19")] == [by_id[i]["label"] for i in ("S18", "S19")] == list(run_labels.ARM_LABELS.values())
    assert short_title(by_id["S2"], center) == "S2 (center recipe, stock DINOv3 init)"
    dirs = [root / by_id[i]["dir"] for i in ("S2", "S18", "S19")]  # a comparison legend without --labels (compare_cyclic_runs.py's own labels)
    assert run_labels.short_run_labels(dirs) == ["S2", *run_labels.ARM_LABELS.values()], run_labels.short_run_labels(dirs)
    assert [arm_colors(by_id)[i] for i in ("S18", "S19")] == list(TAB20[16:18]), arm_colors(by_id)  # the next free colours, in the table's order
    assert list(with_extras(SWEEP, by_id)).index("S18") == SWEEP.index("S16") + 1  # after S16, before the SGD arms


def selftest_seed(tmp):
    """The sweep table's seed, eval_every and eval_cycle0 columns: a non-numeric id (S2b, seed 7), its run directory, labels and colour, its place in the workbook, the deck,
    the ledger, the sweep-wide figures and the AdamW comparison (the fixed presentation order; the table's file order only colours the extras); an extra that is not an adamw arm;
    a table without the columns."""
    from openpyxl import load_workbook

    cols = "id opt lr wd layer_decay patch_embed_mult warmup cycles model init weights normalization seed eval_every eval_cycle0".split()
    stock, constlr = "/w/dinov3_vitb16_pretrain_x.pth", "/w/constlr_final.pth"
    row = lambda i, opt, w, seed, every, c0: [i, opt, "7e-5", "0.01", "0.75", "1.0", "2", "4", "vit_base_dinov3", "dinov3", w, "imagenet", seed, every, c0]

    def mirror(name, rows, live, old=False):  # live = [(id, seed)]: the runs that have a directory with evaluation CSVs
        root, container = Path(tmp) / name, Path(tmp) / f"{name}-container"
        (container / "sweep").mkdir(parents=True)
        table = [cols[:12] if old else cols] + [r[:12] if old else r for r in rows]
        (container / "sweep" / "wave1.tsv").write_text("\n".join("\t".join(r) for r in table) + "\n")
        for ri, (i, seed) in enumerate(live):
            write_eval(root, f"all4_cyclic_dinov3_sweep_{i}_seed_{seed}", ri, (2, 4))
        return root, container

    rows_ = [row("S2b", "adamw", stock, "7", "1", "true"), row("S10", "momentum", stock, "", "", ""), row("S99", "momentum", stock, "100", "2", "false"),
             row("S2", "adamw", stock, "100", "2", "false"), row("S16", "adamw", constlr, "100", "2", "false"), row("S2c", "adamw", stock, "8", "1", "true")]  # file order: neither numeric nor SWEEP order
    live = [("S2b", "7"), ("S10", "100"), ("S99", "100"), ("S2", "100"), ("S16", "100"), ("S2c", "8")]
    shown = [("S2", "100"), ("S16", "100"), ("S2b", "7"), ("S2c", "8"), ("S10", "100"), ("S99", "100")]  # the presentation order: SWEEP, the extra adamw arms after S16, an extra that is not adamw last
    root, container = mirror("seed", rows_, live)
    WARNINGS.clear()
    runs = run_table(root, container)
    by_id = {r["id"]: r for r in runs}
    assert [r["id"] for r in runs] == [i for i, _ in shown], [r["id"] for r in runs]  # the presentation order, not the file's (and S2b would not even sort numerically)
    assert [r["dir"] for r in runs] == [f"all4_cyclic_dinov3_sweep_{i}_seed_{s}" for i, s in shown], runs  # the seed comes from the row; an empty cell is seed 100
    assert tuple(by_id["S2b"]["recipe"][k] for k in SEED_COLS) == ("7", "1", "true"), by_id["S2b"]["recipe"]
    assert by_id["S2b"]["label"] == "S2b (adamw lr 7e-5, wd 0.01, ld 0.75, pe 1.0, init stock, seed 7)" and by_id["S2"]["label"] == "S2 (adamw lr 7e-5, wd 0.01, ld 0.75, pe 1.0, init stock)", (by_id["S2b"]["label"], by_id["S2"]["label"])
    center = by_id["S2"]["recipe"]
    assert [short_title(by_id[i], center) for i in ("S2b", "S2c", "S2", "S16", "S10")] == ["S2b (seed 7)", "S2c (seed 8)", "S2 (center recipe, stock DINOv3 init)", "S16 (init DINO-X ConstLR e30)", "S10 (momentum)"]
    c5 = {**center, "seed": "5"}
    assert short_title({"id": "S2", "label": "", "recipe": c5}, c5) == "S2 (center recipe, stock DINOv3 init, seed 5)"  # the centre arm names a seed other than 100 too
    k = SWEEP.index("S16") + 1
    assert with_extras(SWEEP, by_id) == (*SWEEP[:k], "S2b", "S2c", *SWEEP[k:]) and with_extras(DECK_ORDER, by_id) == (*DECK_ORDER[:k], "S2b", "S2c", *DECK_ORDER[k:]), with_extras(SWEEP, by_id)  # after S16, before the SGD arms, in table order; S99 is not an adamw arm
    assert with_extras(SWEEP, {}) == SWEEP and with_extras(COMPARE["adamw"][1], {}) == COMPARE["adamw"][1]  # no extras: the fixed order
    groups = compare_groups(by_id)
    assert groups["adamw"] == (COMPARE["adamw"][0], (*COMPARE["adamw"][1], "S2b", "S2c")) and groups["init"] == COMPARE["init"] and groups["sgd_frozen"] == COMPARE["sgd_frozen"], groups
    assert any(w.startswith("S99:") and "not adamw" in w for w in WARNINGS) and not any(w.startswith(("S2b:", "S2c:")) for w in WARNINGS), WARNINGS  # an extra that is not adamw is left out, and the build says so
    colors = arm_colors(by_id)
    assert (colors["S2b"], colors["S99"], colors["S2c"]) == (TAB20[16], TAB20[17], TAB20[18]) and all(colors[i] == ARM_COLORS[i] for i in ("S2", "S10", "S16")), colors  # tab20 16, 17, 18, ... in the table's file order (S99 is a table row too)
    assert not {TAB20[14], TAB20[15]} & {colors[i] for i in ("S2b", "S99", "S2c")}  # not the grays
    assert len({colors["S2b"], colors["S2c"], *ARM_COLORS.values()}) == len(ARM_COLORS) + 2  # distinct from every fixed arm
    root2, container2 = mirror("seed2", rows_, [("S2c", "8"), ("S2", "100")])
    by2 = {r["id"]: r for r in run_table(root2, container2)}
    assert arm_colors(by2)["S2c"] == TAB20[18] and with_extras(SWEEP, by2)[k] == "S2c", arm_colors(by2)  # counted over every table row: S2b and S99 have no run here, and S2c keeps its colour
    many = [row(f"X{n}", "adamw", stock, "100", "2", "false") for n in range(1, 7)]
    root5, container5 = mirror("seed-many", many, [(f"X{n}", "100") for n in range(1, 7)])
    assert [arm_colors({r["id"]: r for r in run_table(root5, container5)})[f"X{n}"] for n in range(1, 7)] == [*TAB20[16:], TAB20[15], TAB20[16]]  # six extras: tab20 16-19, the pale gray 15 last, then the sequence starts again

    for i, s in (("S10", "100"), ("S99", "100"), ("S2b", "7"), ("S2", "100")):  # offline ten-crop files, written in the table's order: the ten-crop sheets and the ledger table follow the presentation order
        with open(root / f"cycle_0002_all4_cyclic_dinov3_sweep_{i}_seed_{s}.csv", "w", newline="") as f:
            w = csv.writer(f)
            w.writerow(["dataset", "model", "metric", "class", "value"])
            w.writerows([ds, model, "mAUC", "", 0.9] for ds in DATASETS for model in ("student", "teacher"))
    out = Path(tmp) / "seed-out"
    build(root, container, out, skip_figs=True)
    wb = load_workbook(out / XLSX_NAME)

    def rows(name):
        ws = wb[name]
        head = [c.value for c in ws[1]]
        return [dict(zip(head, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]

    assert [c.value for c in wb["runs"][1]][-3:] == ["seed", "eval_every", "eval_cycle0"], [c.value for c in wb["runs"][1]]
    assert [(r["id"], r["seed"], r["eval_every"], r["eval_cycle0"]) for r in rows("runs")] == [
        ("S2", 100, 2, "false"), ("S16", 100, 2, "false"), ("S2b", 7, 1, "true"), ("S2c", 8, 1, "true"), ("S10", 100, None, None), ("S99", 100, 2, "false")], rows("runs")  # the presentation order, the table's values
    assert [n for n in wb.sheetnames if n in dict(shown)] == [i for i, _ in shown], wb.sheetnames  # the per-run sheets too
    first = lambda name: list(dict.fromkeys(r["run_id"] for r in rows(name)))
    assert first("cyclic_mAUC") == [i for i, _ in shown] and first("offline_eval") == first("tencrop_vs_center") == ["S2", "S2b", "S10", "S99"], (first("cyclic_mAUC"), first("offline_eval"))  # every sheet follows the presentation order
    frag = (out / "tables.html").read_text(encoding="utf-8")
    assert re.findall(r'<td class="id">(S\w+)</td>', frag[frag.index('id="ft-tencrop"'):]) == ["S2", "S2b", "S10"], "ledger"  # the ledger's ten-crop table too, without S99 (not adamw: its offline rows are in the workbook only)
    assert td("id", "S2b") in frag and td("id", "S2c") in frag and td("id", "S99") not in frag, "ledger"  # the extra adamw arms are in the ledger tables
    assert [r["Run"] for r in rows("Last evaluated cycle")] == [i for i, _ in shown if i != "S99"] and "S99" in [r["id"] for r in rows("summary_last")], "ledger sheets"  # the ledger tables' runs in their order; S99 (not adamw) is in the summary sheets only
    tabs = {r["id"]: cycle_table(files) for r in runs if (files := load_eval(root, r))}
    picks, ten = {i: pick(t) for i, t in tabs.items()}, tencrop(root, runs)
    per_run = [t.split(" — ")[0] for t, _ in deck_entries(Path("figs"), tabs, by_id, []) if t.endswith("focused and unfocused — Student and Teacher")]
    assert per_run == [n for n in ("S2 (center recipe, stock DINOv3 init)", "S16 (init DINO-X ConstLR e30)", "S2b (seed 7)", "S2c (seed 8)", "S10 (momentum)") for _ in range(2)], per_run  # the extra arms after S16, before the SGD arms
    sent, real_run = [], run_script
    globals()["run_script"] = lambda script, *a: sent.append(a)
    try:
        compare_figs(root, by_id, set(by_id), out / "figs", Path(tmp))
    finally:
        globals()["run_script"] = real_run
    assert len(sent) == 3, sent  # the AdamW, initialisation and SGD/frozen groups each have two runs
    args, = [a for a in sent if any("S2b_seed_7" in str(x) for x in a)]  # only the AdamW group holds the extra arms
    i, j = args.index("--colors"), args.index("--labels")
    d = "all4_cyclic_dinov3_sweep_"
    assert [x.name for x in args[i - 4 : i]] == [f"{d}S16_seed_100", f"{d}S2b_seed_7", f"{d}S2c_seed_8", f"{d}S2_seed_100"], args  # the table's arms in the AdamW order, the centre last
    assert args[i + 1 : j] == (f"{d}S16_seed_100={ARM_COLORS['S16']}", f"{d}S2b_seed_7={TAB20[16]}", f"{d}S2c_seed_8={TAB20[18]}", f"{d}S2_seed_100={ARM_COLORS['S2']}"), args  # each in its own colour
    assert args[j + 1 :] == (f"{d}S16_seed_100=S16 (init DINO-X ConstLR e30)", f"{d}S2b_seed_7=S2b (seed 7)", f"{d}S2c_seed_8=S2c (seed 8)", f"{d}S2_seed_100=S2 (center recipe, stock DINOv3 init)"), args
    drawn, real_save = {}, save_fig

    def spy(plt, fig, path):
        drawn[path.name] = [ln.get_color() for ln in fig.axes[0].lines]
        real_save(plt, fig, path)

    globals()["save_fig"] = spy
    try:
        sweep_all_figs(ten, tabs, picks, by_id, Path(tmp) / "seed-sw")
    finally:
        globals()["save_fig"] = real_save
    assert drawn["center_MIMIC.png"] == [ARM_COLORS["S2"], ARM_COLORS["S16"], TAB20[16], TAB20[18], ARM_COLORS["S10"]], drawn  # the extra arms are in the sweep-wide figure, each in its own colour

    root4, container4 = mirror("seed-old", [row("S2", "adamw", stock, "", "", "")], [("S2", "100")], old=True)  # a 12-column table, as before the three columns
    old = {r["id"]: r for r in run_table(root4, container4)}
    assert old["S2"]["dir"] == "all4_cyclic_dinov3_sweep_S2_seed_100" and old["S2"]["recipe"]["seed"] == "100" and "seed" not in old["S2"]["label"], old["S2"]
    build(root4, container4, Path(tmp) / "seed-old-out", skip_figs=True)
    wb = load_workbook(Path(tmp) / "seed-old-out" / XLSX_NAME)
    assert [(r["id"], r["seed"], r["eval_every"], r["eval_cycle0"]) for r in rows("runs")] == [("S2", 100, None, None)], rows("runs")  # the default seed; the other two empty


def selftest():
    from openpyxl import load_workbook

    with tempfile.TemporaryDirectory(dir=SCRATCH if SCRATCH.is_dir() else None) as tmp:
        root, container, out = Path(tmp) / "mirror", Path(tmp) / "container", Path(tmp) / "report"
        synth_mirror(root, container)
        ssl_csv = root / "ssl_dinov3_cxr_main_seed_100" / "ssl_losses.csv"
        xlsx = out / XLSX_NAME
        build(root, container, out, skip_figs=True)
        wb = load_workbook(xlsx)

        def rows(name):
            ws = wb[name]
            head = [c.value for c in ws[1]]
            return [dict(zip(head, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]

        def grid(name):  # (header, rows) by cell position: "Best per dataset (teacher)" has four columns named cycle
            ws = wb[name]
            return [c.value for c in ws[1]], [[c.value for c in r] for r in ws.iter_rows(min_row=2)]

        def warned(path, *words):
            return any(str(path) in w and all(x in w for x in words) for w in WARNINGS)

        def retable(path, edit):
            with open(path, newline="") as f:
                table = list(csv.reader(f))
            edit(table)
            with open(path, "w", newline="") as f:
                csv.writer(f).writerows(table)

        assert wb.sheetnames[:8] == ["README", "runs", "Last evaluated cycle", "Best cycle by mean teacher", "Best per dataset (teacher)", "cyclic_mAUC", "cyclic_classAUC", "validation_loss"], wb.sheetnames
        assert len(rows("cyclic_mAUC")) == 2 * 2 * 2 * 4 * 4, len(rows("cyclic_mAUC"))  # runs x splits x cycles x points x datasets
        last = {r["id"]: r["last_cycle"] for r in rows("summary_last")}
        best = {r["id"]: r["best_cycle"] for r in rows("summary_best_mean")}
        assert last == {"S2": 4, "D-long": 10} and best == {"S2": 2, "D-long": 5}, (last, best)
        s2 = next(r for r in rows("runs") if r["id"] == "S2")
        assert (s2["status"], s2["init"], s2["init_tag"]) == ("complete", "dinov3", "stock"), s2  # init is the TSV's own value
        c2, x100 = next(r for r in rows("S2") if r["cycle"] == 2), lambda *a: round(synth(*a) * 100, 4)
        assert c2["T_MIMIC_focused"] == x100(0, 2, 2, 2, 1), c2  # focused = the after_MIMIC row
        assert c2["T_MIMIC_end"] == x100(0, 2, 3, 2, 1), c2  # end of cycle = the after_CheXpert row
        assert c2["T_CheXpert_focused"] == c2["T_CheXpert_end"] == x100(0, 2, 3, 3, 1), c2
        doc = json.loads((root / "results" / "stock_dinov3" / "MIMIC.json").read_text())
        pr = {(r["tag"], r["dataset"]): r for r in rows("probes")}
        assert pr[("stock_dinov3", "MIMIC")]["test_mAUC_mean"] == round(doc["test_mAUC_mean"] * 100, 4)
        assert pr[("stock_dinov3", "ChestXray14")]["overlap_flag"] is True and pr[("stock_dinov3", "MIMIC")]["overlap_flag"] is False
        assert len(rows("offline_eval")) == 8 and {r["protocol"] for r in rows("offline_eval")} == {"tencrop"} and "tencrop" not in wb.sheetnames, wb.sheetnames  # the sheet is renamed
        tv = rows("tencrop_vs_center")
        assert [(r["run_id"], r["cycle"], r["dataset"]) for r in tv] == [("S2", 2, ds) for ds in DATASETS], tv  # one row per dataset
        assert {r["center_source"] for r in tv} == {"in-training"} and all(r["cycle0"] is None for r in rows("runs")), tv  # no offline center crop, no cycle-0 rows on this mirror
        for di, r in enumerate(tv):  # ten-crop mAUC is 0.9 for both models; the center values are S2's cycle-2 end-of-cycle ones (the after_CheXpert rows)
            for k, m in enumerate(("student", "teacher")):
                center = pct(synth(0, 2, 3, di, k))
                assert (r[f"center_{m}"], r[f"tencrop_{m}"], r[f"delta_{m}"]) == (center, 90.0, round(90.0 - center, 2)), r
        assert [(r["epoch"], r["steps"]) for r in rows("ssl_losses_epoch")] == [(0, 3), (1, 3)], rows("ssl_losses_epoch")  # the 1-step trailing epoch 2 is gone
        assert warned(ssl_csv, "epoch 2", "dropped"), WARNINGS
        assert [w for w in WARNINGS if "raydino" in w] == [f"probe tag raydino: no results under {root / 'results'}"], WARNINGS  # a tag that is not on disk: the one existing warning, nothing breaks
        assert "pretraining_overlap_train" in [c.value for c in wb["probes"][1]] and all(r["pretraining_overlap_train"] is None for r in rows("probes")), rows("probes")[:1]  # empty without the key
        ws, hd = wb["probes"], [c.value for c in wb["probes"][1]]
        assert all((ws.cell(2, j).number_format == AUC_FMT) == h.endswith(("_mean", "_std")) for j, h in enumerate(hd, 1)), hd  # number formats follow the header names
        frag = (out / "tables.html").read_text(encoding="utf-8")
        for h in ("Last evaluated cycle", "Best cycle by mean teacher mAUC", "Best teacher value per dataset, with the cycle it occurred",
                  "Linear-probe results by checkpoint", "Our frozen probes (test mAUC, percent, probe s.d. over 10 fits in parentheses)",
                  "Ten-crop vs center crop (test mAUC, percent, teacher; end of cycle)"):
            assert f"<h3>{h}</h3>" in frag, h
        cells = "".join(f'<td class="{"r sep" if di == k == 0 else "r"}">{v}</td>' for di in range(4) for k, v in enumerate((f1(synth(0, 2, 3, di, 1)), "90.0")))
        assert 'id="ft-tencrop"' in frag and f'<tr><td class="id">S2</td><td class="cyc">c2</td>{cells}</tr>' in frag, frag  # percent shown once: no x100 slip, center before ten-crop
        assert "CXR14" not in frag and "Teacher ChestX-ray14" in frag and "MIMIC-CXR center" in frag, frag  # table headers use the display names, not the column ids

        # the ledger sheets: the PDF ledger's three results tables as numbers (S2 = ri 0, cycles 2 and 4, best 2; D-long = ri 1, cycles 5 and 10, best 5; end of cycle = point 3; teacher k = 1, student k = 0)
        v = lambda ri, c: [round(synth(ri, c, 3, di, k) * 100, 1) for k in (1, 0) for di in range(4)]  # the value cells: teacher block, then student block
        d4 = lambda ri, b, e: [round(x - y, 1) for x, y in zip(v(ri, b)[:4], v(ri, e)[:4])]  # best minus last, teacher
        lh = ["Run", "Cycle", "Teacher VinDr-CXR", "Teacher ChestX-ray14", "Teacher MIMIC-CXR", "Teacher CheXpert", "Student VinDr-CXR", "Student ChestX-ray14", "Student MIMIC-CXR", "Student CheXpert", "Status"]
        dh = ["Δ best − last Teacher VinDr-CXR", "Δ best − last Teacher ChestX-ray14", "Δ best − last Teacher MIMIC-CXR", "Δ best − last Teacher CheXpert"]
        h, g = grid("Last evaluated cycle")
        assert h == lh and g == [["D-long", 10, *v(1, 10), None], ["S2", 4, *v(0, 4), "complete"]], (h, g)  # the ledger's runs in its order; the early arm has no status
        h, g = grid("Best cycle by mean teacher")
        assert h == [*lh, *dh, "Last cycle"] and g == [["D-long", 5, *v(1, 5), None, *d4(1, 5, 10), 10], ["S2", 2, *v(0, 2), "complete", *d4(0, 2, 4), 4]], (h, g)
        assert d4(1, 5, 10) == [0.0] * 4 and d4(0, 2, 4) == [10.0] * 4  # D-long's cycles are alike, S2's cycle 2 is 10 points above its cycle 4: the deltas are not vacuous
        h, g = grid("Best per dataset (teacher)")
        assert h == ["Run", "VinDr-CXR best", "cycle", "ChestX-ray14 best", "cycle", "MIMIC-CXR best", "cycle", "CheXpert best", "cycle"], h
        assert g == [["D-long", *[y for x in v(1, 5)[:4] for y in (x, 5)]], ["S2", *[y for x in v(0, 2)[:4] for y in (x, 2)]]], g
        assert {r[0]: r[2::2] for r in g} == {r["id"]: [r[f"{SHORT[d]}_cycle"] for d in DATASETS] for r in rows("summary_best_per_dataset")}  # the cycles are summary_best_per_dataset's
        ok = lambda n, cols: all(c.data_type == "n" and c.number_format == "0.0" for j in cols for (c,) in wb[n].iter_rows(min_row=2, min_col=j, max_col=j))  # numbers (openpyxl reads a whole float back as an int), one decimal
        assert ok("Last evaluated cycle", range(3, 11)) and ok("Best cycle by mean teacher", [*range(3, 11), *range(12, 16)]) and ok("Best per dataset (teacher)", (2, 4, 6, 8))
        assert all(wb[n].cell(2, j).number_format == "General" for n, j in (("Last evaluated cycle", 2), ("Last evaluated cycle", 11), ("Best cycle by mean teacher", 2), ("Best cycle by mean teacher", 11),
                                                                            ("Best cycle by mean teacher", 16), ("Best per dataset (teacher)", 3), ("Best per dataset (teacher)", 9)))  # cycle and status columns have no format
        assert p1(None) is None and (p1(0.1234), p1(0.1236)) == (12.3, 12.4) and all(f"{p1(x):.1f}" == f1(x) for x in (0.0, 0.12345, 0.6255, 0.9995, 1.0)), "p1 is the number f1 prints"
        tb = {"X": {2: {"epoch": 1, "end": {d: (0.12349 + 0.1 * n, 0.98765) for n, d in enumerate(DATASETS)}, "foc": {}}}}  # raw fractions with more than one decimal in percent: the sheet holds the one-decimal number
        assert ledger_sheets(tb, {"X": pick(tb["X"])}, ["X"], {"X": None})[0][2] == [["X", 2, 12.3, 22.3, 32.3, 42.3, 98.8, 98.8, 98.8, 98.8, None]]
        top = [r["README"] for r in rows("README")]
        assert top[0].startswith("DINO-X — numbers workbook (") and top[1] == "Project: DINO-X. Earlier files were titled Step VIII (Demo 1); same data.", top[:2]
        assert [sum(l.startswith(s) for l in top) for s in ("Last evaluated cycle:", "Best cycle by mean teacher:", "Best per dataset (teacher):")] == [1, 1, 1] and not any("Step VIII" in l for l in top[:1] + top[2:]), top
        # the build's self-check: silent on the written sheets, loud for each way a ledger sheet can drift from the summary sheet it re-presents
        sheets, ledger = [(n, *grid(n), {}) for n in wb.sheetnames], ["D-long", "S2"]
        runs = run_table(root, container)
        tabs = {r["id"]: cycle_table(files) for r in runs if (files := load_eval(root, r))}  # the raw fractions the ledger cells round from, read now: the mirror files change below
        assert ledger_disagreements(sheets, ledger, tabs) == [], ledger_disagreements(sheets, ledger, tabs)
        cell = lambda n, i, j: next(rs for m, _, rs, _ in sheets if m == n)[i][j]

        def tampered(name, i, j, to=None):  # the sheets with cell (row i, column j) of the named sheet set to `to`; row i dropped when j is None
            t = [(n, h, [list(r) for r in rs], f) for n, h, rs, f in sheets]
            rs = next(rs for n, _, rs, _ in t if n == name)
            if j is None:
                del rs[i]
            else:
                rs[i][j] = to
            return ledger_disagreements(t, ledger, tabs)

        for case in (("Last evaluated cycle", 0, 2, cell("Last evaluated cycle", 0, 2) + 0.1),  # a value moved by 0.1, in each of the three sheets
                     ("Best cycle by mean teacher", 1, 9, cell("Best cycle by mean teacher", 1, 9) + 0.1),
                     ("Best per dataset (teacher)", 1, 1, cell("Best per dataset (teacher)", 1, 1) + 0.1),
                     ("Last evaluated cycle", 1, 1, 5), ("Best cycle by mean teacher", 0, 1, 6), ("Best per dataset (teacher)", 1, 2, 3),  # a cycle changed
                     ("Last evaluated cycle", 0, None), ("Best cycle by mean teacher", 1, None), ("Best per dataset (teacher)", 0, None),  # a row dropped
                     ("Last evaluated cycle", 1, 10, "in progress (c4 of 4)"), ("Best cycle by mean teacher", 0, 10, "complete"),  # a status changed
                     ("Best cycle by mean teacher", 1, 11, cell("Best cycle by mean teacher", 1, 11) + 0.1), ("Best cycle by mean teacher", 0, 15, 9)):  # a delta changed; the last cycle changed
            assert tampered(*case), case

        # a one-decimal midpoint: a raw 12.34996 % is 12.35 in the summary and 12.3 on the ledger, so the summary's tolerance alone also passes the wrong neighbour 12.4; the raw fraction decides
        from copy import deepcopy

        raw, mt, ms = 0.1234996, deepcopy(tabs), deepcopy(sheets)
        assert (pct(raw), p1(raw)) == (12.35, 12.3)
        mt["D-long"][10]["end"]["VinDrCXR"] = (raw, mt["D-long"][10]["end"]["VinDrCXR"][1])  # D-long's last cycle, VinDr teacher
        row = {n: next(r for r in rs if r[0] == "D-long") for n, _, rs, _ in ms if n in ("summary_last", "Last evaluated cycle", "Best cycle by mean teacher")}
        row["summary_last"][3] = pct(raw)

        def ledger_last(x):  # the ledger's cell for that value set to x, and the delta that follows it
            row["Last evaluated cycle"][2] = x
            row["Best cycle by mean teacher"][11] = round(row["Best cycle by mean teacher"][2] - x, 1)
            return ledger_disagreements(ms, ledger, mt)

        assert ledger_last(p1(raw)) == [], ledger_last(p1(raw))  # the consistent ledger passes
        got = ledger_last(round(p1(raw) + 0.1, 1))  # 12.4, 0.05 from the summary's 12.35
        assert len(got) == 1 and got[0].startswith("Last evaluated cycle: D-long Teacher VinDr-CXR"), got

        bad = root / "all4_cyclic_dinov3_frozen_long_seed_100" / "evaluation" / "CheXpert" / "test_performance.csv"

        def blank_student(table):
            for r in table:
                if r[:4] == ["5", "4", "after_CheXpert", "mAUC"]:
                    r[4] = ""  # one student value becomes unreadable: empty cell plus a warning naming the file, no guess

        def blank_lr(table):
            for r in table[1:]:
                if r[0] == "5":
                    r[table[0].index("lr")] = ""  # the last row of epoch 1: an epoch-end field, counted like the averaged ones

        with open(root / "cycle_0003_all4_cyclic_dinov3_sweep_S2_seed_100.csv", "w", newline="") as f:  # a ten-crop row at a cycle S2 has no evaluation for
            csv.writer(f).writerows([["dataset", "model", "metric", "class", "value"], ["MIMIC", "teacher", "mAUC", "", 0.8]])
        retable(bad, blank_student)
        retable(ssl_csv, blank_lr)
        build(root, container, out, skip_figs=True)
        wb = load_workbook(xlsx)
        row5 = dict(zip([c.value for c in wb["D-long"][1]], [c.value for c in wb["D-long"][2]]))
        assert row5["cycle"] == 5 and row5["S_CheXpert_end"] is None and row5["T_CheXpert_end"] is not None, row5
        assert {r[0].value: r[2].value for r in wb["summary_best_mean"].iter_rows(min_row=2)}["D-long"] == 10  # the incomplete cycle 5 is left out
        assert warned(bad, "unreadable"), WARNINGS
        assert next(r for r in rows("ssl_losses_epoch") if r["epoch"] == 1)["lr_end"] is None and warned(ssl_csv, "unreadable"), WARNINGS
        tv = rows("tencrop_vs_center")
        c3 = [r for r in tv if r["cycle"] == 3]
        assert len(tv) == 5 and len(c3) == 1 and (c3[0]["run_id"], c3[0]["dataset"], c3[0]["tencrop_teacher"]) == ("S2", "MIMIC", 80.0), tv
        assert all(c3[0][k] is None for k in ("center_teacher", "delta_teacher", "center_student", "tencrop_student", "delta_student")), c3  # empty cells, no guess
        assert sum("ten-crop cycle" in w for w in WARNINGS) == 1 and any(all(x in w for x in ("S2", "cycle 3", "MIMIC", "teacher")) and "student" not in w for w in WARNINGS), WARNINGS  # only the model with a ten-crop row is named
        # the sweep figures: a full build also runs the per-run figure scripts (slow), so rebuild build()'s inputs and draw only these
        runs = run_table(root, container)
        tabs = {r["id"]: cycle_table(files) for r in runs if (files := load_eval(root, r))}
        picks, ten = {i: pick(t) for i, t in tabs.items()}, tencrop(root, runs)
        assert sweep_series(ten, tabs, picks, "MIMIC") == ({"S2": [(2, 90.0), (3, 80.0)]}, {"S2": [(c, pct(synth(0, c, 3, 2, 1))) for c in (2, 4)]}, {})  # cycle 3: no center value; no offline center crop: nothing hollow
        tc, cc, _ = sweep_series(ten, {**tabs, "S1": tabs["S2"]}, {**picks, "S1": picks["S2"]}, "MIMIC")  # S1 stands in for a second evaluated sweep arm
        assert list(tc) == ["S2"] and list(cc) == ["S2", "S1"], (tc, cc)  # the center graph shows every evaluated arm, the ten-crop graph only those with ten-crop data
        by_id = {r["id"]: r for r in runs}
        figs = out / "figs" / "sweep_all"
        lines, xlim, real_save = {}, {}, save_fig

        def spy(plt, fig, path):  # the colour of every line drawn and the x range, per figure
            ax = fig.axes[0]
            lines[path.name] = [ln.get_color() for ln in ax.lines]
            xlim[path.name] = tuple(ax.get_xlim())
            real_save(plt, fig, path)

        globals()["save_fig"] = spy
        try:
            sweep_all_figs(ten, tabs, picks, by_id, figs)
        finally:
            globals()["save_fig"] = real_save
        assert lines["center_CheXpert.png"] == lines["tencrop_MIMIC.png"] == [ARM_COLORS["S2"]], lines  # one palette for every arm, in every figure
        assert (xlim["center_CheXpert.png"], xlim["tencrop_MIMIC.png"], xlim["tencrop_VinDrCXR.png"]) == ((2, 4), (2, 3), (2, 3)), xlim  # first to last drawn cycle, no margin either side; a one-point series (tencrop_VinDrCXR: cycle 2 only) gets [c, c + 1]
        assert all((figs / f"{kind}_{ds}.png").is_file() for kind in ("tencrop", "center") for ds in DATASETS), sorted(p.name for p in figs.glob("*"))
        WARNINGS.clear()
        sweep_all_figs([], tabs, picks, by_id, out / "nofigs")  # no ten-crop rows: the center graphs are still drawn, each ten-crop graph is skipped with a warning
        assert len(WARNINGS) == len(DATASETS) and sorted(q.name for q in (out / "nofigs").glob("*")) == sorted(f"center_{ds}.png" for ds in DATASETS), WARNINGS
        WARNINGS.clear()
        sweep_all_figs([], {}, {}, by_id, out / "nothing")  # nothing evaluated either: all eight graphs are skipped
        assert len(WARNINGS) == 2 * len(DATASETS) and not (out / "nothing").exists(), WARNINGS

        def blank_teacher(table):
            for r in table:
                if r[:4] == ["2", "1", "after_CheXpert", "mAUC"]:
                    r[5] = ""  # a center-crop teacher value that cycle 2's ten-crop row can no longer be paired with

        retable(root / "all4_cyclic_dinov3_sweep_S2_seed_100" / "evaluation" / "ChestXray14" / "test_performance.csv", blank_teacher)
        gone = root / "all4_cyclic_dinov3_sweep_S2_seed_100" / "evaluation" / "MIMIC" / "val_performance.csv"
        gone.unlink()  # a live run whose MIMIC validation CSV is not there yet
        build(root, container, out, skip_figs=True)
        wb = load_workbook(xlsx)
        r2 = next(r for r in rows("tencrop_vs_center") if (r["cycle"], r["dataset"]) == (2, "ChestXray14"))
        assert (r2["center_teacher"], r2["delta_teacher"], r2["tencrop_teacher"]) == (None, None, 90.0), r2  # the unreadable center value leaves its cells empty
        assert (r2["center_student"], r2["delta_student"]) == (pct(synth(0, 2, 3, 1, 0)), round(90.0 - pct(synth(0, 2, 3, 1, 0)), 2)), r2  # the other model keeps its values
        assert any(all(x in w for x in ("S2", "cycle 2", "ChestXray14", "teacher")) and "student" not in w for w in WARNINGS), WARNINGS
        s2 = next(r for r in rows("runs") if r["id"] == "S2")
        assert s2["status"] == "in progress" and s2["last_cycle"] == 4 and "S2" in wb.sheetnames, s2  # last cycle still comes from the test rows
        assert {r[0]: r[10] for r in grid("Last evaluated cycle")[1]}["S2"] == "in progress (c4 of 4)", grid("Last evaluated cycle")  # the ledger names the cycle of an unfinished run
        n = Counter((r["run_id"], r["split"]) for r in rows("cyclic_mAUC"))
        assert n[("S2", "test")] == 2 * 4 * 4 and n[("S2", "validation")] == 2 * 4 * 3, n  # every file that exists is kept
        assert warned(gone, "missing") and not any(w.startswith("S2: no evaluation yet") for w in WARNINGS), WARNINGS

        # deck layout v3: the README lines, the comparison call, the 2x2 slide and the probe-page tag choice
        readme = [r["README"] for r in rows("README")]
        assert any(l.startswith("Graphs start at cycle 2:") for l in readme) and any(l.startswith("Comparison and sweep-wide figures draw the teacher only") for l in readme), readme
        sent = []
        real, globals()["run_script"] = run_script, lambda script, *args: sent.append((script, args))
        try:
            compare_figs(root, by_id, {"S2", "D-long"}, out / "figs", Path(tmp))  # in the synthetic mirror only the SGD/frozen group has two runs
        finally:
            globals()["run_script"] = real
        (script, args), = sent
        assert script == "compare_cyclic_runs.py" and ("--models", "teacher") in zip(args, args[1:]) and "--full-length" in args, args
        i, j = args.index("--colors"), args.index("--labels")  # both given after the run directories, as the script asks
        assert [d.name for d in args[i - 2 : i]] == ["all4_cyclic_dinov3_frozen_long_seed_100", "all4_cyclic_dinov3_sweep_S2_seed_100"], args  # the centre arm is passed last
        assert args[i + 1 : j] == (f"all4_cyclic_dinov3_frozen_long_seed_100={ARM_COLORS['D-long']}", f"all4_cyclic_dinov3_sweep_S2_seed_100={ARM_COLORS['S2']}"), args  # every run, in the shared palette
        assert args[j + 1 :] == ("all4_cyclic_dinov3_frozen_long_seed_100=D-long (frozen control, 50 cycles)", "all4_cyclic_dinov3_sweep_S2_seed_100=S2 (center recipe, stock DINOv3 init)"), args  # a run without a recipe: its run-table label
        assert all(rid in ARM_COLORS for _, ids in COMPARE.values() for rid in ids)  # every arm of every comparison group has a palette entry
        from PIL import Image
        from pptx import Presentation
        from pptx.util import Inches

        cells = [Path(tmp) / f"cell{i}.png" for i in range(4)]
        for i, q in enumerate(cells):
            Image.new("RGB", (1100, 620), (60 * i, 120, 200)).save(q)  # the figures' 11 x 6.2 aspect, one colour per dataset
        deck = Path(tmp) / "deck.pptx"
        assert make_deck([("2x2", cells), ("one", cells[0]), ("hole", [cells[0], cells[1], Path(tmp) / "no.png", cells[3]])], deck) == 3
        prs = Presentation(deck)
        pics = [[sh for sh in s.shapes if sh.shape_type == 13] for s in prs.slides]
        assert [len(p) for p in pics] == [4, 1, 3], [len(p) for p in pics]  # a missing PNG leaves its cell empty
        assert [p.image.blob for p in pics[0]] == [q.read_bytes() for q in cells]  # list order = grid order
        tl, tr, bl, br = pics[0]
        assert tl.top == tr.top and bl.top == br.top and tl.left == bl.left and tr.left == br.left and tr.left >= tl.left + tl.width and bl.top >= tl.top + tl.height + Inches(0.1) - 2, pics[0]
        assert all(Inches(0.4) <= p.left and p.left + p.width <= prs.slide_width - Inches(0.4) and Inches(0.85) <= p.top and p.top + p.height <= prs.slide_height - Inches(0.3) for p in pics[0])
        assert abs(pics[1][0].left + pics[1][0].width / 2 - prs.slide_width / 2) <= 2 and pics[1][0].width > 2 * tl.width * 0.9, pics[1]  # a lone PNG still fills the box
        assert any(str(Path(tmp) / "no.png") in w for w in WARNINGS), WARNINGS
        pp = dict.fromkeys(("rad_dino", "dinov2", "stock_dinov3", "stageA_e0010", "stageA_final", "constlr_e0010", "constlr_e0030", "constlr_e0040"))
        assert probe_page_tags(pp) == ["rad_dino", "dinov2", "stock_dinov3", "stageA_final", "constlr_e0040"], probe_page_tags(pp)  # one export per SSL run: the latest
        assert probe_page_tags({k: v for k, v in pp.items() if k != "constlr_e0040"})[-1] == "constlr_e0030"
        assert probe_page_tags(dict.fromkeys(("constlr_e0100", "constlr_e0040"))) == ["constlr_e0100"] and probe_page_tags({"stageA_e0010": None}) == []  # numeric order; no page tag, no pages

        # official names: README, palette, initialisation tags, a long deck title
        assert any(l.startswith("Dataset names: VinDrCXR = VinDr-CXR (column id VinDr)") for l in readme) and sum("Stage A" in l for l in readme) == 1, readme
        import matplotlib
        from matplotlib.colors import to_hex

        assert ARM_COLORS == dict(zip((*SWEEP, "B", "C", "D-long"), map(to_hex, matplotlib.colormaps["tab20"].colors))), ARM_COLORS  # tab20 entries 0..14 in SWEEP order, then B, C, D-long
        assert TAB20 == tuple(map(to_hex, matplotlib.colormaps["tab20"].colors)), TAB20  # the literal is matplotlib's tab20
        cols = "id opt lr wd layer_decay patch_embed_mult warmup cycles model init weights normalization".split()
        m3, c3 = Path(tmp) / "mirror3", Path(tmp) / "container3"
        (c3 / "sweep").mkdir(parents=True)
        tsv = [cols] + [[i, "adamw", "7e-5", "0.01", "0.75", "1.0", "2", "4", "vit_base_dinov3", "dinov3", w, "imagenet"]
                        for i, w in (("S2", "/w/dinov3_vitb16_pretrain_x.pth"), ("S13", "STAGEA_FINAL"), ("S16", "/w/constlr_final.pth"))]
        (c3 / "sweep" / "wave1.tsv").write_text("\n".join("\t".join(r) for r in tsv) + "\n")
        for i in ("S2", "S13", "S16"):
            (m3 / f"all4_cyclic_dinov3_sweep_{i}_seed_100").mkdir(parents=True)
        r3 = {r["id"]: r for r in run_table(m3, c3)}
        assert [r3[i]["recipe"]["init_tag"] for i in ("S2", "S13", "S16")] == ["stock", "DINO-X Cosine e30", "DINO-X ConstLR e30"], r3
        assert [short_title(r3[i], r3["S2"]["recipe"]) for i in ("S13", "S16")] == ["S13 (init DINO-X Cosine e30)", "S16 (init DINO-X ConstLR e30)"]
        assert short_title(r3["S2"], r3["S2"]["recipe"]) == "S2 (center recipe, stock DINOv3 init)" and short_title(r3["S13"], r3["S13"]["recipe"]) == "S13 (center recipe, DINO-X Cosine e30 init)"  # derived from the init tag
        sent2, real_run2 = [], run_script
        globals()["run_script"] = lambda script, *a: sent2.append(a)
        try:
            compare_figs(m3, r3, {"S2", "S13", "S16"}, out / "figs", Path(tmp))  # the AdamW group and the initialisation group both hold these three arms
        finally:
            globals()["run_script"] = real_run2
        want = ("all4_cyclic_dinov3_sweep_S13_seed_100=S13 (init DINO-X Cosine e30)", "all4_cyclic_dinov3_sweep_S16_seed_100=S16 (init DINO-X ConstLR e30)",
                "all4_cyclic_dinov3_sweep_S2_seed_100=S2 (center recipe, stock DINOv3 init)")
        assert len(sent2) == 2 and all(a[a.index("--labels") + 1 :] == want for a in sent2), sent2  # comparison legends say what each curve is initialised from
        make_deck([("t" * 100, cells[0]), ("short", cells[0])], Path(tmp) / "deck2.pptx")
        frames = [next(sh for sh in s.shapes if sh.has_text_frame).text_frame for s in Presentation(Path(tmp) / "deck2.pptx").slides]
        assert [(f.word_wrap, f.paragraphs[0].runs[0].font.size.pt) for f in frames] == [(True, 16.0), (False, 20.0)]  # a long title wraps instead of running off the slide
        WARNINGS.clear()
        make_deck([("t" * 100, cells[0]), ("t" * 111, cells[0])], Path(tmp) / "deck3.pptx")
        assert len(WARNINGS) == 1 and "(111 chars)" in WARNINGS[0], WARNINGS  # 16 pt holds a title on one line up to ~110 characters: only the longer one warns
        titles = [t for t, _ in deck_entries(Path("figs"), tabs, by_id, [("VinDrCXR", Path("p1.png")), ("Summary", Path("p2.png"))])]
        assert titles[:3] == ["Linear probes — VinDr-CXR", "Linear probes — Summary", "DINO-X probe trajectory"], titles[:3]
        assert "DINO-X pretraining loss — dino_global" in titles and "DINO-X learning rate (end of epoch)" in titles
        assert [t for t in titles if t.endswith("focused and unfocused — Student and Teacher")] == [
            f"{n} — {sp} — focused and unfocused — Student and Teacher" for n in ("S2 (center recipe, stock DINOv3 init)", "D-long (frozen control, 50 cycles)") for sp in ("test", "validation")], titles
        assert [t for t in titles if t.startswith("Comparison,")] == [f"Comparison, {g} — {TITLE_DS[ds]} — test — Teacher"
                                                                     for g in ("AdamW arms", "initialisation arms", "SGD arms and frozen control") for ds in DATASETS], titles
        assert [t for t in titles if "sweep arms" in t] == [f"{k}, sweep arms — {TITLE_DS[ds]} — test — Teacher" for k in ("Ten-crop", "Center crop") for ds in DATASETS], titles
        assert max(map(len, titles)) <= 110, titles  # no title reaches the warning

        # the constant-LR run's exports are admitted as found, labelled by tag_label and drawn from the tags present; no KeyError for a tag PROBE_TAGS does not list
        import plot_linear_probes as P

        mir = Path(tmp) / "mirror4"
        P.make_synthetic(str(mir / "results"), tags=("stock_dinov3", "constlr_e0030", "constlr_e0040"))
        both = load_probes(mir)
        assert sorted(both) == ["constlr_e0030", "constlr_e0040", "stock_dinov3"], sorted(both)
        assert [(t, tag_label(t)) for t in probe_page_tags(both)] == [("stock_dinov3", "DINOv3 ViT-B/16 stock (Meta, LVD-1689M)"), ("constlr_e0040", "DINO-X ConstLR e40")]
        assert [t for t, _ in ckpt_rows(both) if t in both] == ["stock_dinov3", "constlr_e0030", "constlr_e0040"] and dict(ckpt_rows(both))["constlr_e0040"] == "DINO-X ConstLR e40"
        assert traj(both)[1] == ("DINO-X ConstLR", (("constlr_e0030", 30), ("constlr_e0040", 40))), traj(both)
        assert {r[0]: r[1] for r in probe_sheets(both)[1]}["constlr_e0040"] == "DINO-X ConstLR e40"
        page = tables_html({}, {}, both, [])
        assert page.count("DINO-X ConstLR e40") == 2 and page.count("DINO-X ConstLR e30") == 2, page  # the checkpoint table and the frozen-probe table
        pages = probe_pages(mir, both, Path(tmp) / "pages", "2026-10-17")
        assert [TITLE_DS.get(n, n) for n, _ in pages] == [*TITLE_DS.values(), "Summary"], pages  # four dataset pages in the one dataset order, then the summary
        fig_trajectory(both, Path(tmp) / "traj.png")
        assert (Path(tmp) / "traj.png").is_file()
        only = {t: d for t, d in both.items() if t != "constlr_e0040"}  # the state today: no e0040 on disk, so nothing about it anywhere
        assert probe_page_tags(only)[-1] == "constlr_e0030" and traj(only)[1][1] == (("constlr_e0030", 30),) and "e40" not in tables_html({}, {}, only, [])
        assert all(tag_label(t) == f"DINO-X ConstLR e{int(t[-4:])}" for t, _ in PROBE_TAGS if CONSTLR.fullmatch(t))  # the table's constant-LR labels are the generated ones
        assert tag_label("rad_dino") == "RAD-DINO ViT-B/14 (Microsoft)" and tag_label("something_else") == "something_else"

        # RayDINO: one extra top-level key in its JSON; it sits after RAD-DINO on the pages and in the tables, its overlap is a sheet column, no dagger
        rmir = Path(tmp) / "mirror5"
        P.make_synthetic(str(rmir / "results"), tags=("rad_dino", "raydino", "dinov2", "stock_dinov3", "stageA_final", "constlr_e0030", "constlr_e0040"))
        train_overlap = "PadChest plus the training splits of CheXpert, MIMIC-CXR and ChestX-ray14"
        for ds in DATASETS:
            f = rmir / "results" / "raydino" / f"{ds}.json"
            f.write_text(json.dumps({**json.loads(f.read_text()), "pretraining_overlap_train": train_overlap}))
        ray = load_probes(rmir)
        assert probe_page_tags(ray) == ["rad_dino", "raydino", "dinov2", "stock_dinov3", "stageA_final", "constlr_e0040"], probe_page_tags(ray)  # RAD-DINO, RayDINO, DINOv2, stock, Cosine e30, latest ConstLR
        assert tag_label("raydino") == "RayDINO ViT-L/16 (Meta, XRay-DINO release)"
        assert any(l == "RayDINO's pretraining corpus is PadChest plus the training splits of CheXpert, MIMIC-CXR and ChestX-ray14 (its Table 2); VinDr-CXR is not in it; "
                   "no test-set overlap, so no dagger; the train-split overlap is recorded in pretraining_overlap_train." for l in readme), readme
        head, rws, _ = probe_sheets(ray)
        col, flag = head.index("pretraining_overlap_train"), head.index("overlap_flag")
        assert col == flag + 1 and {r[0]: r[col] for r in rws} == {t: (train_overlap if t == "raydino" else None) for t in ray}, head[:12]
        assert [r[flag] for r in rws if r[0] == "raydino"] == [False] * 4 and any(r[flag] for r in rws if r[0] == "rad_dino")  # no test-set overlap for RayDINO: no dagger
        assert [t for t, _ in ckpt_rows(ray)][-2:] == ["rad_dino", "raydino"] and dict(ckpt_rows(ray))["raydino"] == "RayDINO (our probe)"
        page = tables_html({}, {}, ray, [])
        assert "RayDINO (our probe)" in page and "RayDINO ViT-L/16 (Meta, XRay-DINO release)" in page, page
        assert len(probe_pages(rmir, ray, Path(tmp) / "pages2", "2026-10-17")) == 5  # plot_linear_probes takes the extra key and the new tag

        selftest_cycle0(tmp)
        selftest_seed(tmp)
        selftest_input_size(tmp)
    print("selftest OK")


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--root", type=Path, default=ROOT, help="mirror root of the Sol run tree")
    ap.add_argument("--container", type=Path, default=CONTAINER, help="container copy (for sweep/wave1.tsv)")
    ap.add_argument("--out-dir", type=Path, default=SCRATCH / "report")
    ap.add_argument("--skip-figs", action="store_true", help="workbook and tables only: no figures, no deck")
    ap.add_argument("--selftest", action="store_true", help="build a tiny synthetic mirror, run the pipeline on it, assert, exit")
    a = ap.parse_args()
    if a.selftest:
        try:
            return selftest()
        except AssertionError as e:
            tb = traceback.extract_tb(e.__traceback__)[-1]
            sys.exit(f"selftest FAILED at line {tb.lineno}: {tb.line} {e}")
    res = build(a.root, a.container, a.out_dir, a.skip_figs)
    pngs = len(list((a.out_dir / "figs").rglob("*.png"))) if (a.out_dir / "figs").is_dir() else 0
    print(f"workbook: {a.out_dir / XLSX_NAME} ({len(res['sheets'])} sheets)")
    print(f"tables:   {a.out_dir / 'tables.html'}")
    print(f"figures:  {pngs} PNG under {a.out_dir / 'figs'}")
    print(f"deck:     {res['slides']} slides" if res["slides"] else "deck:     not built (--skip-figs)")
    print(f"run time: {res['seconds']:.1f} s")
    print(f"WARNINGS ({len(WARNINGS)}):")
    for w in WARNINGS:
        print("  - " + w)


if __name__ == "__main__":
    main()
